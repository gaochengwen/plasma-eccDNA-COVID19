#!/usr/bin/env python3
"""Test the eccDNA burden result against the newly tabulated sequencing covariates.

Adding read pairs, mapping rate and duplicate rate to Supplementary Table S2
(Reviewer 1 minor 3; Reviewer 2 major 3) makes two group differences visible that
were not previously reportable: COVID-19 libraries are deeper and align at a
higher rate than the healthy-control libraries. Any reader can now compute this
from the table, so the manuscript has to say whether the burden result survives
adjustment for them rather than leave the question open.

Writes revise/source_data/seq_qc_covariate_models.tsv and a JSON summary whose
values are quoted verbatim in the Results, Methods and response letters.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
import openpyxl
import statsmodels.api as sm
from scipy.stats import mannwhitneyu

ROOT = Path("/home/gao/eccDNA")
XLSX = ROOT / "Submit" / "Supplementary_Tables_revised.xlsx"
SRC = ROOT / "revise" / "source_data"


def cliffs_delta(a: np.ndarray, b: np.ndarray) -> float:
    return float(np.sign(a[:, None] - b[None, :]).sum() / (len(a) * len(b)))


def main() -> int:
    qc = {r["sample"]: r for r in csv.DictReader(
        (SRC / "sequencing_qc_per_sample.tsv").open(), delimiter="\t")}
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    s2 = {r[0]: {"group": r[1], "count": r[3], "epm": r[4]}
          for r in wb["Table_S2"].iter_rows(min_row=3, values_only=True) if r and r[0]}
    s1 = {r[0]: {"rca": r[1], "sex": r[2], "age": r[3]}
          for r in wb["Table_S1"].iter_rows(min_row=3, values_only=True) if r and r[0]}

    samples = [s for s in s2 if s in qc and s in s1]
    if len(samples) != 78:
        print(f"expected 78 samples, got {len(samples)}")
        return 1
    group = np.array([1 if s2[s]["group"] == "COVID-19" else 0 for s in samples])

    metrics = {
        "read_pairs_millions": np.array([int(qc[s]["read_pairs"]) / 1e6 for s in samples]),
        "mapping_rate_pct": np.array([float(qc[s]["mapping_rate_pct"]) for s in samples]),
        "duplicate_rate_pct": np.array([float(qc[s]["duplicate_rate_pct"]) for s in samples]),
    }

    rows, summary = [], {}
    for name, v in metrics.items():
        a, b = v[group == 1], v[group == 0]
        u, p = mannwhitneyu(a, b, alternative="two-sided")
        d = cliffs_delta(a, b)
        summary[name] = {"covid_median": float(np.median(a)), "hc_median": float(np.median(b)),
                         "U": float(u), "P": float(p), "cliffs_delta": d}
        rows.append(["group_comparison", name, f"{np.median(a):.4f}", f"{np.median(b):.4f}",
                     f"{u:.1f}", f"{p:.4g}", f"{d:+.4f}", "", ""])

    y = np.log2(np.array([s2[s]["epm"] for s in samples], float) + 1)
    design = {
        "group": group.astype(float),
        "log2_rca": np.log2(np.array([s1[s]["rca"] for s in samples], float)),
        "age": np.array([s1[s]["age"] for s in samples], float),
        "male": np.array([1.0 if s1[s]["sex"] == "Male" else 0.0 for s in samples]),
        "mapping_rate_pct": metrics["mapping_rate_pct"],
        "duplicate_rate_pct": metrics["duplicate_rate_pct"],
        "log2_read_pairs": np.log2(np.array([int(qc[s]["read_pairs"]) for s in samples], float)),
    }
    base = ["group", "log2_rca", "age", "male"]
    models = {
        "published": base,
        "plus_mapping_rate": base + ["mapping_rate_pct"],
        "plus_duplicate_rate": base + ["duplicate_rate_pct"],
        "plus_read_pairs": base + ["log2_read_pairs"],
        "plus_all_three": base + ["mapping_rate_pct", "duplicate_rate_pct", "log2_read_pairs"],
    }
    for name, cols in models.items():
        X = sm.add_constant(np.column_stack([design[c] for c in cols]))
        fit = sm.OLS(y, X).fit(cov_type="HC3")
        lo, hi = fit.conf_int()[1]
        summary[name] = {"ratio": float(2 ** fit.params[1]), "ci_low": float(2 ** lo),
                         "ci_high": float(2 ** hi), "P": float(fit.pvalues[1]),
                         "covariates": cols[1:]}
        rows.append(["hc3_model", name, f"{2**fit.params[1]:.4f}", f"{2**lo:.4f}",
                     f"{2**hi:.4f}", f"{fit.pvalues[1]:.4g}", "", "+".join(cols[1:]), "n=78"])
        print(f"  {name:22s} ratio={2**fit.params[1]:5.2f} "
              f"(95% CI {2**lo:.2f}-{2**hi:.2f})  P={fit.pvalues[1]:.3g}")

    SRC.mkdir(parents=True, exist_ok=True)
    with (SRC / "seq_qc_covariate_models.tsv").open("w") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["analysis", "term", "value1", "value2", "value3", "P",
                    "cliffs_delta", "covariates", "note"])
        w.writerows(rows)
    (SRC / "seq_qc_covariate_summary.json").write_text(json.dumps(summary, indent=1))
    print(f"\nwrote {SRC/'seq_qc_covariate_models.tsv'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
