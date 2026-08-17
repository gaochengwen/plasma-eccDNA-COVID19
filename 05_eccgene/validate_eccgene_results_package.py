#!/usr/bin/env python3
"""Validate the organized eccGene statistics, figures, and workbook."""

from __future__ import annotations

import json
import math
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import fitz
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
ANALYSIS = ROOT / "analyse" / "eccGene"
RESULTS = ANALYSIS / "results"
MATRICES = ANALYSIS / "matrices"
REVISE = ROOT / "revise"
SOURCE_DATA = REVISE / "source_data"
REPORT = REVISE / "eccGene_results_validation.txt"

checks: list[str] = []
failures: list[str] = []


def check(condition: bool, message: str) -> None:
    if condition:
        checks.append(f"PASS: {message}")
    else:
        checks.append(f"FAIL: {message}")
        failures.append(message)


def bh_adjust(p_values: np.ndarray) -> np.ndarray:
    order = np.argsort(p_values, kind="mergesort")
    ranked = p_values[order]
    adjusted = ranked * len(ranked) / np.arange(1, len(ranked) + 1)
    adjusted = np.minimum.accumulate(adjusted[::-1])[::-1]
    adjusted = np.clip(adjusted, 0, 1)
    result = np.empty_like(adjusted)
    result[order] = adjusted
    return result


def fisher_two_sided(table: list[list[int]]) -> float:
    """Exact two-sided Fisher p-value for a 2 x 2 table.

    This is the standard fixed-margin probability-ordering definition used by
    SciPy's ``fisher_exact(..., alternative="two-sided")``.
    """
    a, b = table[0]
    c, d = table[1]
    row_1 = a + b
    column_1 = a + c
    total = a + b + c + d
    denominator = math.comb(total, row_1)

    def probability(x: int) -> float:
        return (
            math.comb(column_1, x)
            * math.comb(total - column_1, row_1 - x)
            / denominator
        )

    lower = max(0, row_1 - (total - column_1))
    upper = min(row_1, column_1)
    observed = probability(a)
    tolerance = observed * 1e-12
    return min(
        1.0,
        sum(
            probability(x)
            for x in range(lower, upper + 1)
            if probability(x) <= observed + tolerance
        ),
    )


def read_tsv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, sep="\t")


def validate_statistics() -> None:
    metadata = read_tsv(ANALYSIS / "sample_metadata.tsv")
    abundance = read_tsv(RESULTS / "junction_abundance_wilcoxon.tsv")
    log1p = read_tsv(RESULTS / "junction_abundance_log1p_sensitivity.tsv")
    detection = read_tsv(RESULTS / "junction_detection_fisher.tsv")
    check(
        metadata["group"].value_counts().to_dict() == {"COVID-19": 39, "HC": 39},
        "sample metadata contains 39 COVID-19 and 39 HC samples",
    )
    check(
        len(abundance) == len(detection) == 26642,
        "abundance and detection analyses each contain 26,642 tested genes",
    )
    check(
        set(abundance["Gene"]) == set(detection["Gene"]),
        "abundance and detection tables use the same gene universe",
    )
    check(
        abundance["Gene"].tolist() == log1p["Gene"].tolist(),
        "raw-EA and log1p(EA) sensitivity tables use the same genes",
    )
    check(
        np.allclose(
            abundance["wilcoxon_p"],
            log1p["wilcoxon_p_log1p"],
            rtol=0,
            atol=1e-14,
        )
        and np.allclose(
            abundance["wilcoxon_q_BH"],
            log1p["wilcoxon_q_BH_log1p"],
            rtol=0,
            atol=1e-14,
        ),
        "Wilcoxon p/q values are invariant to the monotonic log1p(EA) transform",
    )
    check(
        np.allclose(
            abundance["wilcoxon_q_BH"],
            bh_adjust(abundance["wilcoxon_p"].to_numpy()),
            rtol=1e-8,
            atol=1e-12,
        ),
        "abundance BH-FDR values independently reproduce from Wilcoxon p-values",
    )
    check(
        np.allclose(
            detection["fisher_q_BH"],
            bh_adjust(detection["fisher_p"].to_numpy()),
            rtol=1e-8,
            atol=1e-12,
        ),
        "detection BH-FDR values independently reproduce from Fisher p-values",
    )

    for gene in ("LINC01255", "PAX1", "CLEC12B", "BCL3", "PROCR"):
        row = detection.loc[detection["Gene"] == gene].iloc[0]
        table = [
            [int(row["covid_detected"]), int(row["covid_not_detected"])],
            [int(row["hc_detected"]), int(row["hc_not_detected"])],
        ]
        p_value = fisher_two_sided(table)
        check(
            math.isclose(p_value, row["fisher_p"], rel_tol=1e-9, abs_tol=1e-15),
            f"{gene} two-sided Fisher p-value independently reproduces",
        )

    check(
        int((abundance["wilcoxon_q_BH"] < 0.05).sum()) == 12793,
        "12,793 genes meet abundance BH q<0.05",
    )
    check(
        int(
            (
                (abundance["wilcoxon_q_BH"] < 0.05)
                & (abundance["log2FC_mean_EA_COVID_vs_HC"].abs() >= 1)
            ).sum()
        )
        == 3402,
        "3,402 abundance genes also meet |mean-EA log2FC|>=1",
    )
    significant_detection = detection.loc[detection["fisher_q_BH"] < 0.05]
    check(
        len(significant_detection) == 20603
        and set(significant_detection["direction"]) == {"COVID_up"},
        "20,603 detection genes meet BH q<0.05 and all are COVID-up",
    )


def validate_matrices() -> None:
    counts = read_tsv(MATRICES / "junction_counts.tsv").set_index("Gene")
    detected = read_tsv(MATRICES / "junction_detected.tsv").set_index("Gene")
    ea = read_tsv(MATRICES / "junction_EA.tsv").set_index("Gene")
    check(
        counts.index.equals(detected.index) and counts.index.equals(ea.index),
        "junction count, detection, and EA matrices have identical gene order",
    )
    check(
        counts.columns.equals(detected.columns) and counts.columns.equals(ea.columns),
        "junction count, detection, and EA matrices have identical sample order",
    )
    check(
        np.array_equal((counts.to_numpy() > 0).astype(np.int8), detected.to_numpy()),
        "binary detection matrix is exactly count>0",
    )
    check(
        np.array_equal(counts.to_numpy() == 0, ea.to_numpy() == 0),
        "EA is zero exactly where the unique-isotype count is zero",
    )
    column_sums = ea.sum(axis=0).to_numpy()
    check(
        np.allclose(column_sums, 1_000_000, rtol=0, atol=0.05),
        "junction EA sums to 1,000,000 within each sample",
    )


def validate_sensitivity_and_figure_sources() -> None:
    stability = read_tsv(RESULTS / "candidate_stability.tsv")
    abundance = read_tsv(RESULTS / "junction_abundance_wilcoxon.tsv")
    detection = read_tsv(RESULTS / "junction_detection_fisher.tsv")
    stable_abundance = stability[
        stability["abundance_significant_same_direction_all_definitions"] == 1
    ]
    stable_detection = stability[
        stability["detection_significant_same_direction_all_definitions"] == 1
    ]
    check(len(stable_abundance) == 3271, "3,271 abundance genes are stable across definitions")
    check(len(stable_detection) == 11707, "11,707 detection genes are stable across definitions")

    a_lookup = abundance.set_index("Gene")
    expected_a = stable_abundance[
        stable_abundance["junction_abundance_log2FC"] >= 1
    ].copy()
    expected_a["cliffs_delta"] = expected_a["Gene"].map(
        a_lookup["cliffs_delta_COVID_vs_HC"]
    )
    expected_a = (
        expected_a[expected_a["cliffs_delta"] > 0]
        .sort_values(
            ["junction_abundance_q", "cliffs_delta", "Gene"],
            ascending=[True, False, True],
        )
        .head(30)["Gene"]
        .tolist()
    )
    observed_a = read_tsv(SOURCE_DATA / "Figure_3A_abundance_genes.tsv")[
        "Gene"
    ].tolist()
    check(
        observed_a == expected_a,
        "Figure 3A genes exactly match the declared stable-abundance selection rule",
    )

    stable_detection_genes = set(stable_detection["Gene"])
    expected_b = (
        detection[
            detection["Gene"].isin(stable_detection_genes)
            & (detection["fisher_q_BH"] < 0.05)
            & (detection["frequency_difference_COVID_minus_HC"] > 0)
        ]
        .sort_values(
            ["fisher_q_BH", "frequency_difference_COVID_minus_HC", "Gene"],
            ascending=[True, False, True],
        )
        .head(15)["Gene"]
        .tolist()
    )
    observed_b = read_tsv(SOURCE_DATA / "Figure_3B_detection_genes.tsv")[
        "Gene"
    ].tolist()
    check(
        observed_b == expected_b,
        "Figure 3B genes exactly match the declared stable-detection selection rule",
    )

    replica_a = read_tsv(
        SOURCE_DATA / "Figure_3A_replica_abundance_genes.tsv"
    )
    expected_replica_a = set(
        abundance.loc[
            (abundance["wilcoxon_q_BH"] < 0.05)
            & (abundance["log2FC_mean_EA_COVID_vs_HC"].abs() >= 1),
            "Gene",
        ]
    )
    check(
        set(replica_a["Gene"]) == expected_replica_a
        and len(replica_a) == 3402,
        "replica Figure 3A contains exactly the 3,402 dual-threshold abundance genes",
    )
    replica_b = read_tsv(
        SOURCE_DATA / "Figure_3B_replica_detection_genes.tsv"
    )
    expected_replica_b_selected = (
        detection[
            detection["Gene"].isin(expected_replica_a)
            & (detection["fisher_q_BH"] < 0.05)
            & (detection["hc_detected"] == 0)
            & (detection["frequency_difference_COVID_minus_HC"] > 0)
        ]
        .sort_values(
            ["fisher_q_BH", "covid_detected", "Gene"],
            ascending=[True, False, True],
            kind="stable",
        )
        .head(30)
    )
    expected_replica_b = (
        expected_replica_b_selected.sort_values(
            ["covid_detected", "fisher_q_BH", "Gene"],
            ascending=[False, True, True],
            kind="stable",
        )["Gene"]
        .tolist()
    )
    check(
        replica_b["Gene"].tolist() == expected_replica_b,
        "replica Figure 3B contains the declared top 30 HC-absent detection genes "
        "ordered by decreasing recurrence",
    )
    replica_b_matrix = read_tsv(
        SOURCE_DATA / "Figure_3B_replica_detection_matrix.tsv"
    ).set_index("Gene")
    covid_matrix_columns = list(replica_b_matrix.columns[:39])
    covid_column_recurrence = replica_b_matrix[covid_matrix_columns].sum(axis=0)
    expected_covid_columns = sorted(
        covid_matrix_columns,
        key=lambda sample: (-int(covid_column_recurrence[sample]), sample),
    )
    check(
        covid_matrix_columns == expected_covid_columns
        and covid_column_recurrence.is_monotonic_decreasing,
        "replica Figure 3B COVID-19 samples are ordered by decreasing displayed-gene "
        "recurrence",
    )

    expected_query = set(
        abundance.loc[
            (abundance["wilcoxon_q_BH"] < 0.05)
            & (abundance["log2FC_mean_EA_COVID_vs_HC"] >= 1)
            & (abundance["cliffs_delta_COVID_vs_HC"] > 0),
            "Gene",
        ]
    )
    observed_query = set(
        (SOURCE_DATA / "eccGene_gProfiler_query_genes.txt")
        .read_text(encoding="utf-8")
        .splitlines()
    )
    expected_background = set(abundance["Gene"])
    observed_background = set(
        (SOURCE_DATA / "eccGene_gProfiler_background_genes.txt")
        .read_text(encoding="utf-8")
        .splitlines()
    )
    check(
        observed_query == expected_query and len(observed_query) == 3094,
        "g:Profiler query is exactly the 3,094 declared COVID-up abundance genes",
    )
    check(
        observed_background == expected_background
        and len(observed_background) == 26642,
        "g:Profiler custom background is exactly all 26,642 tested genes",
    )
    response = json.loads(
        (SOURCE_DATA / "eccGene_gProfiler_response.json").read_text(encoding="utf-8")
    )
    all_terms = read_tsv(SOURCE_DATA / "eccGene_gProfiler_all_significant_terms.tsv")
    check(
        len(response["result"]) == len(all_terms) == 55
        and bool((all_terms["p_value"] < 0.05).all()),
        "g:Profiler response and exported table contain the same 55 significant terms",
    )

    circles = read_tsv(SOURCE_DATA / "Figure_3D_sorted_recurrent_circles.tsv")
    sample_columns = [column for column in circles if column not in {"eccDNA", "recurrence"}]
    check(
        len(circles) == 27
        and len(sample_columns) == 39
        and bool((circles["recurrence"] >= 10).all()),
        "Figure 3D contains 27 intervals recurring in >=10 of 39 COVID-19 samples",
    )
    check(
        np.array_equal(
            circles["recurrence"].to_numpy(),
            circles[sample_columns].sum(axis=1).to_numpy(),
        ),
        "Figure 3D recurrence values equal the row sums of the binary matrix",
    )


def validate_workbook() -> None:
    workbook_path = REVISE / "Supplementary_Tables_revised.xlsx"
    with zipfile.ZipFile(workbook_path) as archive:
        check(archive.testzip() is None, "supplementary workbook ZIP members pass CRC checks")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    expected_dimensions = {
        "Table_S1": (80, 17),
        "Table_S2": (80, 5),
        "Table_S3": (80, 26),
        "Table_S4": (18, 79),
        "Table_S5": (548, 8),
        "Table_S6": (548, 10),
        "Table_S7": (3404, 83),
        "Table_S8": (69, 27),
        "Table_S9": (2611, 16),
        "Table_S10": (826, 16),
        "Table_S11": (174, 19),
    }
    observed = {}
    table_s7_title = None
    table_s7_header = None
    for sheet in workbook.worksheets:
        sheet.calculate_dimension(force=True)
        observed[sheet.title] = (sheet.max_row, sheet.max_column)
        if sheet.title == "Table_S7":
            table_s7_title = next(
                sheet.iter_rows(min_row=1, max_row=1, values_only=True)
            )
            table_s7_header = next(
                sheet.iter_rows(min_row=2, max_row=2, values_only=True)
            )
    workbook.close()
    check(
        observed == expected_dimensions,
        "supplementary workbook sheet order and dimensions match the expected package",
    )
    check(
        table_s7_title is not None
        and table_s7_title[0]
        == "Supplementary Table S7 EccDNA abundance values in eccGenes with "
        "significant differences (BH-FDR < 0.05 and absolute log2FC >= 1)",
        "Table S7 has the requested abundance-only title",
    )
    check(
        table_s7_header is not None
        and table_s7_header[0] == "Gene"
        and table_s7_header[-4:]
        == (
            "log2FC",
            "Cliffs_delta",
            "p_wilcoxon",
            "adj_wilcoxon_BH",
        ),
        "Table S7 contains 78 sample EA columns and only abundance statistics",
    )


def pdf_page_size(path: Path) -> tuple[float, float]:
    with fitz.open(path) as document:
        if len(document) != 1:
            raise ValueError(f"{path.name} must contain exactly one page")
        rectangle = document[0].rect
    return rectangle.width, rectangle.height


def validate_figures() -> None:
    expected = {
        "Figure_3_AB": ((518.74, 221.102), (2161, 921)),
        "Figure_S3": ((518.74, 260.787), (2161, 1086)),
    }
    for stem, (expected_pdf, expected_png) in expected.items():
        pdf = REVISE / f"{stem}.pdf"
        svg = REVISE / f"{stem}.svg"
        png = REVISE / f"{stem}.png"
        width, height = pdf_page_size(pdf)
        check(
            math.isclose(width, expected_pdf[0], abs_tol=0.02)
            and math.isclose(height, expected_pdf[1], abs_tol=0.02),
            f"{stem} PDF has the declared 183-mm-wide page size",
        )
        with Image.open(png) as image:
            check(
                image.size == expected_png,
                f"{stem} PNG has the expected 300-dpi pixel dimensions",
            )
        root = ET.parse(svg).getroot()
        svg_width = float(root.attrib["width"].removesuffix("pt"))
        check(
            math.isclose(svg_width, 518.74, abs_tol=0.02),
            f"{stem} SVG has the declared 183-mm width",
        )
        if stem == "Figure_3_AB":
            with fitz.open(pdf) as document:
                image_sizes = {
                    (image[2], image[3])
                    for image in document[0].get_images(full=True)
                }
            check(
                len(image_sizes) == 2
                and (510, 64) in image_sizes
                and any(
                    width in {1115, 1116} and height == 1244
                    for width, height in image_sizes
                ),
                "Figure_3_AB PDF rasterizes only panel A and its colorbar; "
                "panel B cells remain vector objects",
            )

    final_figure = REVISE / "Figure 3.pdf"
    width, height = pdf_page_size(final_figure)
    check(
        math.isclose(width, 518.74, abs_tol=0.02)
        and math.isclose(height, 709.333, abs_tol=0.02),
        "final Figure 3 PDF has the declared 183-mm-wide page size",
    )


def main() -> None:
    validate_statistics()
    validate_matrices()
    validate_sensitivity_and_figure_sources()
    validate_workbook()
    validate_figures()
    REPORT.write_text("\n".join(checks) + "\n", encoding="utf-8")
    print("\n".join(checks))
    if failures:
        raise SystemExit(f"{len(failures)} validation checks failed")


if __name__ == "__main__":
    main()
