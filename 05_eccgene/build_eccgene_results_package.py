#!/usr/bin/env python3
"""Build the eccGene revision result summary and supplementary workbook."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path("/home/gao/eccDNA")
ANALYSIS = ROOT / "analyse" / "eccGene"
RESULTS = ANALYSIS / "results"
MATRICES = ANALYSIS / "matrices"
REVISE = ROOT / "revise"
SOURCE_DATA = REVISE / "source_data"
WORKBOOK = REVISE / "Supplementary_Tables_revised.xlsx"
SUMMARY = REVISE / "eccGene_results_summary.md"
MANIFEST = REVISE / "eccGene_results_manifest.tsv"

TITLE_FILL = PatternFill("solid", fgColor="3C5488")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
WHITE_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Arial", size=8, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=8, color="000000")
SECTION_FONT = Font(name="Arial", size=9, bold=True, color="1F1F1F")


def read_tsv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, sep="\t")


def scalar(value: str):
    if value == "":
        return None
    lowered = value.lower()
    if lowered in {"inf", "+inf", "infinity", "+infinity"}:
        return "Inf"
    if lowered in {"-inf", "-infinity"}:
        return "-Inf"
    try:
        integer = int(value)
        if str(integer) == value:
            return integer
    except ValueError:
        pass
    try:
        number = float(value)
        if math.isfinite(number):
            return number
    except ValueError:
        pass
    return value


def number_format(header: str) -> str:
    key = header.lower()
    if any(token in key for token in ("_p", "_q", "p_value", "fdr")):
        return "0.00E+00"
    if any(
        token in key
        for token in (
            "mean",
            "median",
            "frequency",
            "log2",
            "cliff",
            "odds_ratio",
            "95ci",
            "jaccard",
            "precision",
            "recall",
            "denominator",
        )
    ):
        return "0.0000"
    return "General"


def style_title(ws, title: str, max_column: int) -> None:
    ws.cell(1, 1, title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_column)
    cell = ws.cell(1, 1)
    cell.fill = TITLE_FILL
    cell.font = TITLE_FONT
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22


def style_header_row(ws, row: int, max_column: int) -> None:
    for cell in ws[row][:max_column]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def set_widths(ws, headers: list[str], observed: list[int]) -> None:
    for index, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 2, observed[index - 1] + 1)
        if header == "Gene":
            width = max(width, 16)
        if any(token in header.lower() for token in ("path", "description", "rule")):
            width = min(max(width, 28), 55)
        else:
            width = min(max(width, 10), 24)
        ws.column_dimensions[get_column_letter(index)].width = width


def add_tsv_sheet(
    workbook,
    index: int,
    name: str,
    title: str,
    path: Path,
    row_filter: Callable[[dict[str, str]], bool] | None = None,
) -> int:
    ws = workbook.create_sheet(name, index)
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        headers = reader.fieldnames
        if not headers:
            raise ValueError(f"No header found in {path}")
        style_title(ws, title, len(headers))
        ws.append(headers)
        style_header_row(ws, 2, len(headers))
        observed = [len(header) for header in headers]
        for raw in reader:
            if row_filter is not None and not row_filter(raw):
                continue
            values = [scalar(raw[header]) for header in headers]
            ws.append(values)
            if ws.max_row <= 202:
                for column, value in enumerate(values):
                    observed[column] = max(
                        observed[column], len(str(value if value is not None else ""))
                    )
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{ws.max_row}"
    set_widths(ws, headers, observed)
    return index + 1


def add_definition_sheet(workbook, index: int) -> int:
    ws = workbook.create_sheet("Table_S7D_Definitions", index)
    style_title(
        ws,
        "Supplementary Table 7D. Locked eccGene annotation, assignment, and "
        "normalization rules",
        9,
    )
    rules = [
        ("Reference assembly", "hg38"),
        (
            "Reference annotation",
            "UCSC refGene table; local source file dated 29 May 2025",
        ),
        ("Reference MD5", "77af23e0a7f318fca6ecf517d4296e05"),
        ("Retained chromosomes", "chr1-chr22, chrX, chrY, chrM"),
        (
            "Transcript handling",
            "All refGene txStart-txEnd spans collapsed to one bounding gene-body "
            "interval per gene symbol per chromosome",
        ),
        (
            "Within-sample de-duplication",
            "Unique (chromosome, start, end) Circle-Map BED coordinates",
        ),
        (
            "Primary definition",
            "The 1-bp BED start interval [start,start+1) lies within a merged gene body",
        ),
        (
            "Sensitivity definition 1",
            "Any full eccDNA interval-gene-body intersection",
        ),
        (
            "Sensitivity definition 2",
            "floor[(start+end)/2] lies within a merged gene body",
        ),
        (
            "Multi-gene assignment",
            "One circle intersecting multiple genes contributes one count to each; "
            "counts are not fractional",
        ),
        (
            "Multiple isoforms",
            "Each distinct per-sample (chromosome,start,end) interval assigned to a "
            "gene contributes one count",
        ),
        (
            "Gene length",
            "Sum of collapsed gene-body interval lengths across retained chromosomes",
        ),
        (
            "EA formula",
            "EA_sg = [(X_sg/L_g) / sum_h(X_sh/L_h)] x 10^6",
        ),
        (
            "Abundance endpoint",
            "Two-sided Wilcoxon rank-sum test on per-sample EA; BH-FDR; mean-EA "
            "log2 fold change with 1e-6 pseudocount; Cliff's delta",
        ),
        (
            "Detection endpoint",
            "detected = 1 when X_sg > 0; two-sided Fisher's exact test; conditional "
            "odds ratio and exact 95% CI; BH-FDR",
        ),
        ("Sample size", "39 COVID-19 and 39 healthy controls"),
    ]
    ws.append(["Rule", "Locked specification"])
    style_header_row(ws, 2, 2)
    for rule, specification in rules:
        ws.append([rule, specification])
        for cell in ws[ws.max_row][:2]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary = read_tsv(RESULTS / "definition_summary.tsv")
    overlap = read_tsv(RESULTS / "definition_overlap_summary.tsv")
    software = read_tsv(RESULTS / "software_versions.tsv")
    sections: list[tuple[str, pd.DataFrame]] = [
        ("Definition-level result summary", summary),
        ("Pairwise significant-set overlap", overlap),
        ("Software versions on QDUH", software),
    ]
    for label, frame in sections:
        ws.append([])
        ws.append([label])
        section_row = ws.max_row
        ws.merge_cells(
            start_row=section_row,
            start_column=1,
            end_row=section_row,
            end_column=max(2, len(frame.columns)),
        )
        ws.cell(section_row, 1).fill = SECTION_FILL
        ws.cell(section_row, 1).font = SECTION_FONT
        ws.append(frame.columns.tolist())
        style_header_row(ws, ws.max_row, len(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            ws.append(list(values))
            for cell in ws[ws.max_row][: len(frame.columns)]:
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top")

    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 31
    ws.column_dimensions["B"].width = 75
    for column in range(3, 10):
        ws.column_dimensions[get_column_letter(column)].width = 18
    return index + 1


def replace_table_s7(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    table_s7_sheets = [
        sheet for sheet in workbook.sheetnames if sheet.startswith("Table_S7")
    ]
    insertion_index = min(
        (workbook.sheetnames.index(sheet) for sheet in table_s7_sheets),
        default=6,
    )
    for sheet in table_s7_sheets:
        workbook.remove(workbook[sheet])

    index = insertion_index
    index = add_tsv_sheet(
        workbook,
        index,
        "Table_S7A_Abundance",
        "Supplementary Table 7A. Junction-based eccGene abundance analysis "
        "(all 26,642 tested genes)",
        RESULTS / "junction_abundance_wilcoxon.tsv",
    )
    index = add_tsv_sheet(
        workbook,
        index,
        "Table_S7B_Detection",
        "Supplementary Table 7B. Junction-based eccGene detection-frequency analysis "
        "(all 26,642 tested genes)",
        RESULTS / "junction_detection_fisher.tsv",
    )
    index = add_tsv_sheet(
        workbook,
        index,
        "Table_S7C_Stability",
        "Supplementary Table 7C. Assignment-stable genes significant in the same "
        "direction under all three definitions for at least one endpoint",
        RESULTS / "candidate_stability.tsv",
        row_filter=lambda row: (
            row["abundance_significant_same_direction_all_definitions"] == "1"
            or row["detection_significant_same_direction_all_definitions"] == "1"
        ),
    )
    index = add_definition_sheet(workbook, index)
    index = add_tsv_sheet(
        workbook,
        index,
        "Table_S7E_Enrichment",
        "Supplementary Table 7E. g:Profiler functional over-representation analysis "
        "(55 significant terms; custom tested-gene background)",
        SOURCE_DATA / "eccGene_gProfiler_all_significant_terms.tsv",
    )
    index = add_tsv_sheet(
        workbook,
        index,
        "Table_S7F_Recurrent",
        "Supplementary Table 7F. Exact eccDNA intervals detected in at least 10 "
        "COVID-19 samples and no healthy-control samples",
        SOURCE_DATA / "Figure_3D_sorted_recurrent_circles.tsv",
    )
    add_tsv_sheet(
        workbook,
        index,
        "Table_S7G_Sample_QC",
        "Supplementary Table 7G. Sample-level eccGene assignment and input quality "
        "metrics",
        RESULTS / "sample_eccgene_metrics.tsv",
    )

    temporary = workbook_path.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, workbook_path)


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
WORKSHEET_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
)
WORKSHEET_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
)


def xml_text(value) -> str:
    text = str(value)
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)
    return escape(text)


def xml_cell(reference: str, value, style: int) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return f'<c r="{reference}" s="{style}" t="b"><v>{int(value)}</v></c>'
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return f'<c r="{reference}" s="{style}"><v>{value}</v></c>'
    return (
        f'<c r="{reference}" s="{style}" t="inlineStr"><is><t>'
        f"{xml_text(value)}</t></is></c>"
    )


def tsv_rows(
    path: Path,
    row_filter: Callable[[dict[str, str]], bool] | None = None,
):
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        headers = reader.fieldnames
        if not headers:
            raise ValueError(f"No header found in {path}")
        for raw in reader:
            if row_filter is None or row_filter(raw):
                yield [scalar(raw[header]) for header in headers]


def tsv_headers(path: Path) -> list[str]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        return next(reader)


def count_tsv_rows(
    path: Path,
    row_filter: Callable[[dict[str, str]], bool] | None = None,
) -> int:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        return sum(1 for row in reader if row_filter is None or row_filter(row))


def definition_rows():
    rules = [
        ["Reference assembly", "hg38"],
        [
            "Reference annotation",
            "UCSC refGene table; local source file dated 29 May 2025",
        ],
        ["Reference MD5", "77af23e0a7f318fca6ecf517d4296e05"],
        ["Retained chromosomes", "chr1-chr22, chrX, chrY, chrM"],
        [
            "Transcript handling",
            "All refGene txStart-txEnd spans collapsed to one bounding gene-body "
            "interval per gene symbol per chromosome",
        ],
        [
            "Within-sample de-duplication",
            "Unique (chromosome, start, end) Circle-Map BED coordinates",
        ],
        [
            "Primary definition",
            "The 1-bp BED start interval [start,start+1) lies within a merged gene body",
        ],
        [
            "Sensitivity definition 1",
            "Any full eccDNA interval-gene-body intersection",
        ],
        [
            "Sensitivity definition 2",
            "floor[(start+end)/2] lies within a merged gene body",
        ],
        [
            "Multi-gene assignment",
            "One circle intersecting multiple genes contributes one count to each; "
            "counts are not fractional",
        ],
        [
            "Multiple isoforms",
            "Each distinct per-sample (chromosome,start,end) interval assigned to a "
            "gene contributes one count",
        ],
        [
            "Gene length",
            "Sum of collapsed gene-body interval lengths across retained chromosomes",
        ],
        ["EA formula", "EA_sg = [(X_sg/L_g) / sum_h(X_sh/L_h)] x 10^6"],
        [
            "Abundance endpoint",
            "Two-sided Wilcoxon rank-sum test on per-sample EA; BH-FDR; mean-EA "
            "log2 fold change with 1e-6 pseudocount; Cliff's delta",
        ],
        [
            "Detection endpoint",
            "detected = 1 when X_sg > 0; two-sided Fisher's exact test; conditional "
            "odds ratio and exact 95% CI; BH-FDR",
        ],
        ["Sample size", "39 COVID-19 and 39 healthy controls"],
    ]
    yield from rules
    for label, path in (
        ("Definition-level result summary", RESULTS / "definition_summary.tsv"),
        (
            "Pairwise significant-set overlap",
            RESULTS / "definition_overlap_summary.tsv",
        ),
        ("Software versions on QDUH", RESULTS / "software_versions.tsv"),
    ):
        yield []
        yield [label]
        headers = tsv_headers(path)
        yield headers
        yield from tsv_rows(path)


def write_streamed_sheet(
    archive: zipfile.ZipFile,
    member: str,
    title: str,
    headers: list[str],
    rows: Iterable[list],
    data_row_count: int,
    max_columns: int | None = None,
) -> None:
    max_columns = max(max_columns or len(headers), len(headers))
    last_column = get_column_letter(max_columns)
    column_xml = []
    for index in range(1, max_columns + 1):
        header = headers[index - 1] if index <= len(headers) else ""
        width = 16
        if header in {"Gene", "eccDNA"}:
            width = 18
        elif any(token in header.lower() for token in ("path", "description", "rule")):
            width = 36
        column_xml.append(
            f'<col min="{index}" max="{index}" width="{width}" customWidth="1"/>'
        )
    prefix = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{MAIN_NS}">'
        f'<dimension ref="A1:{last_column}{data_row_count + 2}"/>'
        '<sheetViews><sheetView workbookViewId="0">'
        '<pane xSplit="1" ySplit="2" topLeftCell="B3" activePane="bottomRight" '
        'state="frozen"/>'
        '<selection pane="bottomRight" activeCell="B3" sqref="B3"/>'
        "</sheetView></sheetViews>"
        '<sheetFormatPr baseColWidth="10" defaultRowHeight="15"/>'
        f"<cols>{''.join(column_xml)}</cols><sheetData>"
    )
    with archive.open(member, "w") as handle:
        handle.write(prefix.encode("utf-8"))
        title_cell = xml_cell("A1", title, 19)
        handle.write(f'<row r="1" ht="22" customHeight="1">{title_cell}</row>'.encode())
        header_cells = "".join(
            xml_cell(f"{get_column_letter(column)}2", value, 12)
            for column, value in enumerate(headers, start=1)
        )
        handle.write(
            f'<row r="2" ht="30" customHeight="1">{header_cells}</row>'.encode()
        )
        row_number = 2
        buffer: list[str] = []
        for values in rows:
            row_number += 1
            cells = "".join(
                xml_cell(f"{get_column_letter(column)}{row_number}", value, 14)
                for column, value in enumerate(values, start=1)
            )
            buffer.append(f'<row r="{row_number}">{cells}</row>')
            if len(buffer) == 1000:
                handle.write("".join(buffer).encode("utf-8"))
                buffer.clear()
        if buffer:
            handle.write("".join(buffer).encode("utf-8"))
        suffix = (
            "</sheetData>"
            f'<autoFilter ref="A2:{get_column_letter(len(headers))}{row_number}"/>'
            f'<mergeCells count="1"><mergeCell ref="A1:{last_column}1"/></mergeCells>'
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" '
            'header="0.3" footer="0.3"/>'
            "</worksheet>"
        )
        handle.write(suffix.encode("utf-8"))


def replace_table_s7_streaming(workbook_path: Path) -> None:
    abundance = RESULTS / "junction_abundance_wilcoxon.tsv"
    detection = RESULTS / "junction_detection_fisher.tsv"
    stability = RESULTS / "candidate_stability.tsv"
    enrichment = SOURCE_DATA / "eccGene_gProfiler_all_significant_terms.tsv"
    recurrent = SOURCE_DATA / "Figure_3D_sorted_recurrent_circles.tsv"
    sample_qc = RESULTS / "sample_eccgene_metrics.tsv"
    stable_filter = lambda row: (
        row["abundance_significant_same_direction_all_definitions"] == "1"
        or row["detection_significant_same_direction_all_definitions"] == "1"
    )
    definition_data = list(definition_rows())
    specs = [
        (
            "Table_S7A_Abundance",
            "Supplementary Table 7A. Junction-based eccGene abundance analysis "
            "(all 26,642 tested genes)",
            tsv_headers(abundance),
            lambda: tsv_rows(abundance),
            count_tsv_rows(abundance),
            None,
        ),
        (
            "Table_S7B_Detection",
            "Supplementary Table 7B. Junction-based eccGene detection-frequency "
            "analysis (all 26,642 tested genes)",
            tsv_headers(detection),
            lambda: tsv_rows(detection),
            count_tsv_rows(detection),
            None,
        ),
        (
            "Table_S7C_Stability",
            "Supplementary Table 7C. Assignment-stable genes significant in the "
            "same direction under all three definitions for at least one endpoint",
            tsv_headers(stability),
            lambda: tsv_rows(stability, stable_filter),
            count_tsv_rows(stability, stable_filter),
            None,
        ),
        (
            "Table_S7D_Definitions",
            "Supplementary Table 7D. Locked eccGene annotation, assignment, "
            "normalization, and analysis rules",
            ["Rule", "Locked specification"],
            lambda: iter(definition_data),
            len(definition_data),
            13,
        ),
        (
            "Table_S7E_Enrichment",
            "Supplementary Table 7E. g:Profiler functional over-representation "
            "analysis (55 significant terms; custom tested-gene background)",
            tsv_headers(enrichment),
            lambda: tsv_rows(enrichment),
            count_tsv_rows(enrichment),
            None,
        ),
        (
            "Table_S7F_Recurrent",
            "Supplementary Table 7F. Exact eccDNA intervals detected in at least "
            "10 COVID-19 samples and no healthy-control samples",
            tsv_headers(recurrent),
            lambda: tsv_rows(recurrent),
            count_tsv_rows(recurrent),
            None,
        ),
        (
            "Table_S7G_Sample_QC",
            "Supplementary Table 7G. Sample-level eccGene assignment and input "
            "quality metrics",
            tsv_headers(sample_qc),
            lambda: tsv_rows(sample_qc),
            count_tsv_rows(sample_qc),
            None,
        ),
    ]

    ET.register_namespace("", MAIN_NS)
    ET.register_namespace("r", OFFICE_REL_NS)
    temporary = workbook_path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(workbook_path, "r") as source:
        workbook_root = ET.fromstring(source.read("xl/workbook.xml"))
        rels_root = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        content_root = ET.fromstring(source.read("[Content_Types].xml"))
        sheets = workbook_root.find(f"{{{MAIN_NS}}}sheets")
        if sheets is None:
            raise ValueError("Workbook has no sheets element")

        rel_by_id = {
            relation.attrib["Id"]: relation
            for relation in rels_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
        }
        removed_members: set[str] = set()
        insertion_index = len(sheets)
        for index, sheet in list(enumerate(list(sheets))):
            if not sheet.attrib.get("name", "").startswith("Table_S7"):
                continue
            insertion_index = min(insertion_index, index)
            relation_id = sheet.attrib[f"{{{OFFICE_REL_NS}}}id"]
            relation = rel_by_id[relation_id]
            target = relation.attrib["Target"].lstrip("/")
            member = target if target.startswith("xl/") else f"xl/{target}"
            removed_members.add(member)
            removed_members.add(
                member.replace("worksheets/", "worksheets/_rels/") + ".rels"
            )
            sheets.remove(sheet)
            rels_root.remove(relation)
            for override in list(
                content_root.findall(f"{{{CONTENT_TYPE_NS}}}Override")
            ):
                if override.attrib.get("PartName") == f"/{member}":
                    content_root.remove(override)

        existing_sheet_numbers = [
            int(match.group(1))
            for name in source.namelist()
            if (match := re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", name))
        ]
        next_sheet_number = max(existing_sheet_numbers, default=0) + 1
        next_sheet_id = (
            max(int(sheet.attrib["sheetId"]) for sheet in sheets) + 1
        )
        relationship_numbers = [
            int(match.group(1))
            for relation in rels_root.findall(
                f"{{{PACKAGE_REL_NS}}}Relationship"
            )
            if (match := re.fullmatch(r"rId(\d+)", relation.attrib["Id"]))
        ]
        next_relationship = max(relationship_numbers, default=0) + 1
        new_members: list[tuple[str, tuple]] = []
        for offset, spec in enumerate(specs):
            member = f"xl/worksheets/sheet{next_sheet_number}.xml"
            relation_id = f"rId{next_relationship}"
            sheet = ET.Element(
                f"{{{MAIN_NS}}}sheet",
                {
                    "name": spec[0],
                    "sheetId": str(next_sheet_id),
                    f"{{{OFFICE_REL_NS}}}id": relation_id,
                },
            )
            sheets.insert(insertion_index + offset, sheet)
            ET.SubElement(
                rels_root,
                f"{{{PACKAGE_REL_NS}}}Relationship",
                {
                    "Id": relation_id,
                    "Type": WORKSHEET_REL_TYPE,
                    "Target": member.removeprefix("xl/"),
                },
            )
            ET.SubElement(
                content_root,
                f"{{{CONTENT_TYPE_NS}}}Override",
                {
                    "PartName": f"/{member}",
                    "ContentType": WORKSHEET_CONTENT_TYPE,
                },
            )
            new_members.append((member, spec))
            next_sheet_number += 1
            next_sheet_id += 1
            next_relationship += 1

        modified_members = {
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "[Content_Types].xml",
        }
        with zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6
        ) as destination:
            for item in source.infolist():
                if item.filename in modified_members or item.filename in removed_members:
                    continue
                destination.writestr(item, source.read(item.filename))
            destination.writestr(
                "xl/workbook.xml",
                ET.tostring(
                    workbook_root,
                    encoding="utf-8",
                    xml_declaration=True,
                ),
            )
            destination.writestr(
                "xl/_rels/workbook.xml.rels",
                ET.tostring(
                    rels_root,
                    encoding="utf-8",
                    xml_declaration=True,
                ),
            )
            destination.writestr(
                "[Content_Types].xml",
                ET.tostring(
                    content_root,
                    encoding="utf-8",
                    xml_declaration=True,
                ),
            )
            for member, spec in new_members:
                _, title, headers, row_factory, data_row_count, max_columns = spec
                write_streamed_sheet(
                    destination,
                    member,
                    title,
                    headers,
                    row_factory(),
                    data_row_count=data_row_count,
                    max_columns=max_columns,
                )

    with zipfile.ZipFile(temporary, "r") as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"Corrupt workbook member: {bad_member}")
    validation = load_workbook(temporary, read_only=True, data_only=False)
    expected = [spec[0] for spec in specs]
    observed = [
        name for name in validation.sheetnames if name.startswith("Table_S7")
    ]
    validation.close()
    if observed != expected:
        raise ValueError(f"Unexpected Table S7 sheet order: {observed}")
    os.replace(temporary, workbook_path)


def styled_write_only_row(ws, values: list, kind: str) -> list[WriteOnlyCell]:
    cells: list[WriteOnlyCell] = []
    for value in values:
        cell = WriteOnlyCell(ws, value=value)
        if kind == "title":
            cell.fill = TITLE_FILL
            cell.font = TITLE_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        elif kind == "header":
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        cells.append(cell)
    return cells


def configure_clean_sheet(ws, headers: list[str], max_columns: int) -> None:
    ws.freeze_panes = "B3"
    for index in range(1, max_columns + 1):
        header = headers[index - 1] if index <= len(headers) else ""
        width = 14
        if header in {"Gene", "eccDNA"}:
            width = 18
        elif any(token in header.lower() for token in ("path", "description", "rule")):
            width = 36
        elif len(header) > 20:
            width = 22
        ws.column_dimensions[get_column_letter(index)].width = width


def append_clean_table(
    ws,
    title: str,
    headers: list[str],
    rows: Iterable[list],
    max_columns: int | None = None,
) -> int:
    max_columns = max(max_columns or len(headers), len(headers))
    configure_clean_sheet(ws, headers, max_columns)
    title_values = [title] + [None] * (max_columns - 1)
    ws.append(styled_write_only_row(ws, title_values, "title"))
    ws.append(
        styled_write_only_row(
            ws, headers + [None] * (max_columns - len(headers)), "header"
        )
    )
    row_count = 2
    for values in rows:
        ws.append(values)
        row_count += 1
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{row_count}"
    return row_count


def previous_table_s7_sample_order() -> list[str]:
    previous = load_workbook(
        ROOT / "Maintext" / "Supplementary_Tables.xlsx",
        read_only=True,
        data_only=True,
    )
    header = list(
        next(
            previous["Table_S7"].iter_rows(
                min_row=2,
                max_row=2,
                values_only=True,
            )
        )
    )
    previous.close()
    samples = [str(value) for value in header[1:-3] if value is not None]
    if len(samples) != 78:
        raise ValueError(f"Expected 78 previous Table S7 samples, observed {len(samples)}")
    return samples


def significant_abundance_matrix_rows(
    sample_order: list[str],
) -> Iterable[list]:
    statistics = read_tsv(RESULTS / "junction_abundance_wilcoxon.tsv")
    statistics = statistics.loc[
        (statistics["wilcoxon_q_BH"] < 0.05)
        & (statistics["log2FC_mean_EA_COVID_vs_HC"].abs() >= 1)
    ].set_index("Gene")
    expected_genes = set(statistics.index)
    observed_genes: set[str] = set()
    with (MATRICES / "junction_EA.tsv").open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        available_samples = set(reader.fieldnames or []) - {"Gene"}
        if set(sample_order) != available_samples:
            raise ValueError("Previous Table S7 sample order does not match the EA matrix")
        for row in reader:
            gene = row["Gene"]
            if gene not in expected_genes:
                continue
            observed_genes.add(gene)
            result = statistics.loc[gene]
            yield [
                gene,
                *[round(float(row[sample]), 4) for sample in sample_order],
                round(float(result["log2FC_mean_EA_COVID_vs_HC"]), 6),
                round(float(result["cliffs_delta_COVID_vs_HC"]), 6),
                float(result["wilcoxon_p"]),
                float(result["wilcoxon_q_BH"]),
            ]
    missing = expected_genes - observed_genes
    if missing:
        raise ValueError(f"Significant genes missing from EA matrix: {len(missing)}")


def clean_s7_specs() -> list[tuple]:
    sample_order = previous_table_s7_sample_order()
    headers = [
        "Gene",
        *sample_order,
        "log2FC",
        "Cliffs_delta",
        "p_wilcoxon",
        "adj_wilcoxon_BH",
    ]
    return [
        (
            "Table_S7",
            "Supplementary Table 7 EccDNA abundance values in eccGenes with "
            "significant differences (BH-FDR < 0.05 and absolute log2FC >= 1)",
            headers,
            lambda: significant_abundance_matrix_rows(sample_order),
            None,
        ),
    ]


def rebuild_supplementary_workbook(workbook_path: Path) -> None:
    source = load_workbook(workbook_path, read_only=True, data_only=False)
    clean = Workbook(write_only=True)
    clean.iso_dates = True
    clean.calculation.fullCalcOnLoad = True
    clean.calculation.forceFullCalc = True
    clean.calculation.calcMode = "auto"

    source_sheets = {
        name: source[name]
        for name in source.sheetnames
        if not name.startswith("Table_S7")
    }
    s7_specs = clean_s7_specs()
    output_order = [
        "Table_S1",
        "Table_S2",
        "Table_S3",
        "Table_S4",
        "Table_S5",
        "Table_S6",
        *[spec[0] for spec in s7_specs],
        "Table_S8",
        "Table_S9",
        "Table_S10",
    ]
    spec_by_name = {spec[0]: spec for spec in s7_specs}

    for name in output_order:
        ws = clean.create_sheet(name)
        if name in spec_by_name:
            _, title, headers, row_factory, max_columns = spec_by_name[name]
            append_clean_table(
                ws,
                title,
                headers,
                row_factory(),
                max_columns=max_columns,
            )
            continue

        old = source_sheets[name]
        iterator = old.iter_rows(values_only=True)
        try:
            title_row = list(next(iterator))
            header_row = list(next(iterator))
        except StopIteration as error:
            raise ValueError(f"Source sheet {name} has fewer than two rows") from error
        max_columns = old.max_column or max(len(title_row), len(header_row))
        headers = [
            str(value) if value is not None else "" for value in header_row[:max_columns]
        ]
        configure_clean_sheet(ws, headers, max_columns)
        ws.append(styled_write_only_row(ws, title_row[:max_columns], "title"))
        ws.append(styled_write_only_row(ws, header_row[:max_columns], "header"))
        row_count = 2
        for row in iterator:
            ws.append(list(row[:max_columns]))
            row_count += 1
        ws.auto_filter.ref = f"A2:{get_column_letter(max_columns)}{row_count}"

    source.close()
    temporary = workbook_path.with_suffix(".clean.tmp.xlsx")
    clean.save(temporary)
    with zipfile.ZipFile(temporary, "r") as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"Corrupt clean workbook member: {bad_member}")
    validation = load_workbook(temporary, read_only=True, data_only=False)
    observed = validation.sheetnames
    validation.close()
    if observed != output_order:
        raise ValueError(f"Unexpected clean workbook sheet order: {observed}")
    os.replace(temporary, workbook_path)


def format_p(value: float) -> str:
    return f"{value:.3g}"


def gene_row(frame: pd.DataFrame, gene: str) -> pd.Series:
    matches = frame.loc[frame["Gene"] == gene]
    if len(matches) != 1:
        raise ValueError(f"Expected one row for {gene}, observed {len(matches)}")
    return matches.iloc[0]


def build_summary() -> None:
    abundance = read_tsv(RESULTS / "junction_abundance_wilcoxon.tsv")
    detection = read_tsv(RESULTS / "junction_detection_fisher.tsv")
    stability = read_tsv(RESULTS / "candidate_stability.tsv")
    overlap = read_tsv(RESULTS / "definition_overlap_summary.tsv")
    enrichment = read_tsv(SOURCE_DATA / "eccGene_gProfiler_all_significant_terms.tsv")
    circles = read_tsv(SOURCE_DATA / "Figure_3D_sorted_recurrent_circles.tsv")

    abundance_sig = abundance["wilcoxon_q_BH"] < 0.05
    high_effect = abundance_sig & (
        abundance["log2FC_mean_EA_COVID_vs_HC"].abs() >= 1
    )
    enrichment_query = (
        abundance_sig
        & (abundance["log2FC_mean_EA_COVID_vs_HC"] >= 1)
        & (abundance["cliffs_delta_COVID_vs_HC"] > 0)
    )
    stable_abundance = int(
        stability["abundance_significant_same_direction_all_definitions"].sum()
    )
    stable_detection = int(
        stability["detection_significant_same_direction_all_definitions"].sum()
    )
    stable_high_effect = int(
        (
            (stability["abundance_significant_same_direction_all_definitions"] == 1)
            & (stability["junction_abundance_log2FC"].abs() >= 1)
        ).sum()
    )

    lines = [
        "# eccGene reanalysis: organized results",
        "",
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
        "",
        "## Locked primary analysis",
        "",
        "- Genome/annotation: hg38 UCSC `refGene`; source MD5 "
        "`77af23e0a7f318fca6ecf517d4296e05`.",
        "- Primary gene assignment: the 1-bp Circle-Map BED start/junction "
        "`[start,start+1)` lies within a merged gene body.",
        "- Sensitivity assignments: any full-interval intersection and midpoint "
        "within a merged gene body.",
        "- Per-sample de-duplication: unique `(chromosome,start,end)` intervals.",
        "- EA: `[(X_sg/L_g) / sum_h(X_sh/L_h)] x 10^6`.",
        "- Abundance: two-sided Wilcoxon rank-sum test, BH-FDR, mean-EA log2 fold "
        "change, and Cliff's delta.",
        "- Detection: binary `X_sg>0`, two-sided Fisher's exact test, conditional "
        "odds ratio with exact 95% CI, and BH-FDR.",
        "",
        "## Main numerical results",
        "",
        f"- Tested genes: {len(abundance):,}.",
        f"- Abundance BH q<0.05: {int(abundance_sig.sum()):,}; among these, "
        f"{int(high_effect.sum()):,} had |mean-EA log2FC|>=1.",
        f"- COVID-up abundance query used for enrichment: "
        f"{int(enrichment_query.sum()):,} genes (q<0.05, log2FC>=1, "
        "positive Cliff's delta).",
        f"- Detection-frequency BH q<0.05: "
        f"{int((detection['fisher_q_BH'] < 0.05).sum()):,}; all significant "
        "directions were COVID-up.",
        f"- Significant in the same direction under all three definitions: "
        f"{stable_abundance:,} abundance genes and {stable_detection:,} detection "
        f"genes; {stable_high_effect:,} abundance genes also had "
        "|junction log2FC|>=1.",
        f"- g:Profiler returned {len(enrichment):,} significant terms using all "
        f"{len(abundance):,} tested genes as the custom background.",
        f"- Recurrent exact intervals: {len(circles):,} detected in >=10 COVID-19 "
        "samples and no healthy-control samples.",
        "",
        "## Representative genes",
        "",
        "| Gene | Abundance result | Detection result | Sensitivity interpretation |",
        "|---|---|---|---|",
    ]
    for gene in ("LINC01255", "PAX1", "CLEC12B", "BCL3", "PROCR"):
        a = gene_row(abundance, gene)
        d = gene_row(detection, gene)
        s = gene_row(stability, gene)
        abundance_text = (
            f"log2FC={a['log2FC_mean_EA_COVID_vs_HC']:.3g}, "
            f"Cliff's delta={a['cliffs_delta_COVID_vs_HC']:.3g}, "
            f"q={format_p(a['wilcoxon_q_BH'])}"
        )
        detection_text = (
            f"{int(d['covid_detected'])}/39 vs {int(d['hc_detected'])}/39; "
            f"OR={d['odds_ratio_conditional_COVID_vs_HC']:.3g}, "
            f"95% CI {d['odds_ratio_95CI_low']:.3g}-"
            f"{d['odds_ratio_95CI_high']:.3g}; "
            f"q={format_p(d['fisher_q_BH'])}"
        )
        stable_a = bool(s["abundance_significant_same_direction_all_definitions"])
        stable_d = bool(s["detection_significant_same_direction_all_definitions"])
        sensitivity = (
            f"abundance stable={'yes' if stable_a else 'no'}; "
            f"detection stable={'yes' if stable_d else 'no'}"
        )
        lines.append(
            f"| {gene} | {abundance_text} | {detection_text} | {sensitivity} |"
        )

    lines.extend(
        [
            "",
            "## Sensitivity and interpretation",
            "",
        ]
    )
    for endpoint in ("abundance_wilcoxon", "detection_fisher"):
        subset = overlap.loc[overlap["endpoint"] == endpoint]
        endpoint_label = "Abundance" if endpoint.startswith("abundance") else "Detection"
        j_i = subset.loc[
            (subset["definition_a"] == "junction")
            & (subset["definition_b"] == "interval"),
            "jaccard",
        ].iloc[0]
        j_m = subset.loc[
            (subset["definition_a"] == "junction")
            & (subset["definition_b"] == "midpoint"),
            "jaccard",
        ].iloc[0]
        lines.append(
            f"- {endpoint_label} significant-set Jaccard: "
            f"junction-midpoint {j_m:.3f}; junction-full interval {j_i:.3f}."
        )
    lines.extend(
        [
            "- The abundance and detection endpoints must remain separate. Fisher's "
            "exact test is used only for binary detection.",
            "- Detection frequency is influenced by the larger total Circle-seq "
            "eccDNA burden in COVID-19 and is not sufficient on its own to establish "
            "disease specificity.",
            "- `eccGene` is a positional assignment label; it does not demonstrate "
            "that an intact or functional gene is carried or regulated by an eccDNA.",
            "- Pathway results are over-representation associations and do not "
            "demonstrate pathway activation.",
            "",
            "## Figure and table mapping",
            "",
            "- `Figure_3.*`: stable abundance heatmap, stable detection-frequency "
            "comparison, functional enrichment, and recurrent exact circles.",
            "- `Figure_S3.*`: cross-definition stable-association counts, pairwise "
            "Jaccard concordance, and effect-size sensitivity for top stable genes.",
            "- `Graphical_abstract.*`: descriptive revised workflow and principal "
            "observations.",
            "- `Supplementary_Tables_revised.xlsx`: Table S7 contains the "
            "junction-based EA matrix for 3,402 genes meeting BH-FDR < 0.05 and "
            "absolute log2FC >= 1, "
            "followed by log2FC, Cliff's delta, Wilcoxon p, and BH q; all other "
            "supplementary tables are retained.",
            "- `source_data/`: plotted data, g:Profiler request/response, query and "
            "background gene lists, and figure run metadata.",
            "",
            "The manuscript and reviewer-response DOCX files were intentionally not "
            "modified in this result-organization step.",
            "",
        ]
    )
    SUMMARY.write_text("\n".join(lines), encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_manifest() -> None:
    files: list[tuple[Path, str]] = [
        (SUMMARY, "Human-readable organized result summary"),
        (
            REVISE / "eccGene_results_validation.txt",
            "Independent statistical and file-integrity validation report",
        ),
        (
            WORKBOOK,
            "Merged supplementary workbook with one revised abundance-only Table S7",
        ),
        (REVISE / "Figure_3.pdf", "Main Figure 3, vector PDF"),
        (REVISE / "Figure_3.svg", "Main Figure 3, editable SVG"),
        (REVISE / "Figure_3.png", "Main Figure 3, 300-dpi PNG"),
        (
            REVISE / "Figure_3_AB.pdf",
            "Revised Figure 3A-B replica, vector PDF",
        ),
        (
            REVISE / "Figure_3_AB.svg",
            "Revised Figure 3A-B replica, editable SVG",
        ),
        (
            REVISE / "Figure_3_AB.png",
            "Revised Figure 3A-B replica, 300-dpi PNG",
        ),
        (REVISE / "Figure_S3.pdf", "Sensitivity figure, vector PDF"),
        (REVISE / "Figure_S3.svg", "Sensitivity figure, editable SVG"),
        (REVISE / "Figure_S3.png", "Sensitivity figure, 300-dpi PNG"),
        (REVISE / "Graphical_abstract.pdf", "Graphical abstract, vector PDF"),
        (REVISE / "Graphical_abstract.svg", "Graphical abstract, editable SVG"),
        (REVISE / "Graphical_abstract.png", "Graphical abstract, 300-dpi PNG"),
    ]
    files.extend(
        (path, "Figure source data or reproducibility metadata")
        for path in sorted(SOURCE_DATA.iterdir())
        if path.is_file()
        and path.name.startswith(
            ("Figure_3", "Figure_S3", "eccGene_", "covid_specific_exact")
        )
    )
    with MANIFEST.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["relative_path", "description", "size_bytes", "sha256"])
        for path, description in files:
            writer.writerow(
                [
                    path.relative_to(REVISE),
                    description,
                    path.stat().st_size,
                    sha256(path),
                ]
            )


def main() -> None:
    rebuild_supplementary_workbook(WORKBOOK)
    build_summary()
    build_manifest()
    print(f"Updated {WORKBOOK}")
    print(f"Created {SUMMARY}")
    print(f"Created {MANIFEST}")


if __name__ == "__main__":
    main()
