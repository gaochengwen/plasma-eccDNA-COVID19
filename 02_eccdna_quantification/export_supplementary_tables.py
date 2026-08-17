#!/usr/bin/env python3
"""Export manuscript supplementary Excel tables to TSV inputs.

This script is intentionally small and deterministic.  The QDUH runtime used
for the revision statistics does not provide pandas/openpyxl, so the Excel
workbook is converted locally to plain TSV files before being copied to QDUH.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def norm_header(value: object) -> str:
    text = "" if value is None else str(value).strip()
    replacements = {
        " ": "_",
        "-": "_",
        "/": "_",
        "(": "",
        ")": "",
        "′": "",
        "'": "",
        "ﬁ": "fi",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_").lower()


def clean_cell(value: object) -> object:
    if value is None:
        return ""
    return value


def write_tsv(path: Path, header: Iterable[str], rows: Iterable[Iterable[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(list(header))
        for row in rows:
            writer.writerow([clean_cell(value) for value in row])


def read_rectangular_sheet(workbook, sheet_name: str, header_row: int = 2) -> tuple[list[str], list[list[object]]]:
    ws = workbook[sheet_name]
    iterator = ws.iter_rows(values_only=True)
    header_values = None
    for row_number, row in enumerate(iterator, start=1):
        if row_number == header_row:
            header_values = list(row)
            break
    if header_values is None:
        raise ValueError(f"Header row {header_row} was not found in {sheet_name}")
    header = [norm_header(value) for value in header_values]
    rows: list[list[object]] = []
    width = len(header)
    for row in iterator:
        row = list(row)
        if len(row) < width:
            row = row + [None] * (width - len(row))
        elif len(row) > width:
            row = row[:width]
        if any(value is not None for value in row):
            rows.append(row)
    return header, rows


def export_long_table(workbook, sheet_name: str, output_path: Path) -> None:
    header, rows = read_rectangular_sheet(workbook, sheet_name)
    write_tsv(output_path, header, rows)


def export_chromosome_distribution(workbook, output_path: Path) -> None:
    header, rows = read_rectangular_sheet(workbook, "Table_S3")
    sample_idx = header.index("sample_id")
    group_idx = header.index("group")
    chromosome_cols = [(idx, name) for idx, name in enumerate(header) if name.startswith("chr")]
    long_rows = []
    for row in rows:
        sample_id = row[sample_idx]
        group = row[group_idx]
        for idx, chromosome in chromosome_cols:
            long_rows.append([sample_id, group, chromosome, row[idx]])
    write_tsv(output_path, ["sample_id", "group", "chromosome", "normalized_fraction_per_mb"], long_rows)


def export_repeat_class_enrichment(workbook, output_path: Path, group_by_sample: dict[str, str]) -> None:
    ws = workbook["Table_S4"]
    iterator = ws.iter_rows(values_only=True)
    header_row = None
    for row_number, row in enumerate(iterator, start=1):
        if row_number == 2:
            header_row = list(row)
            break
    if header_row is None:
        raise ValueError("Header row 2 was not found in Table_S4")
    sample_ids = [str(value).strip() for value in header_row[1:] if value is not None]
    rows = []
    for row in iterator:
        repeat_class = row[0] if row else None
        if repeat_class is None:
            continue
        values = list(row[1:])
        for offset, sample_id in enumerate(sample_ids):
            rows.append(
                [
                    sample_id,
                    group_by_sample.get(sample_id, ""),
                    repeat_class,
                    values[offset] if offset < len(values) else "",
                ]
            )
    write_tsv(output_path, ["sample_id", "group", "repeat_class", "normalized_mapping_ratio"], rows)


def build_group_map(workbook) -> dict[str, str]:
    header, rows = read_rectangular_sheet(workbook, "Table_S2")
    sample_idx = header.index("sample_id")
    group_idx = header.index("group")
    return {str(row[sample_idx]).strip(): str(row[group_idx]).strip() for row in rows}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    output_dir = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    group_by_sample = build_group_map(workbook)

    export_long_table(workbook, "Table_S1", output_dir / "sample_metadata.tsv")
    export_long_table(workbook, "Table_S2", output_dir / "sample_burden.tsv")
    export_chromosome_distribution(workbook, output_dir / "chromosome_distribution.tsv")
    export_repeat_class_enrichment(workbook, output_dir / "repeat_class_enrichment.tsv", group_by_sample)
    export_long_table(workbook, "Table_S5", output_dir / "gene_feature_characteristics.tsv")
    export_long_table(workbook, "Table_S6", output_dir / "gene_feature_enrichment_scores.tsv")
    export_long_table(workbook, "Table_S8", output_dir / "histone_encode_overlap.tsv")
    export_long_table(workbook, "Table_S9", output_dir / "histone_sars_cov2_overlap.tsv")

    manifest_rows = []
    for path in sorted(output_dir.glob("*.tsv")):
        stat = path.stat()
        manifest_rows.append([path.name, stat.st_size, int(stat.st_mtime)])
    write_tsv(output_dir / "export_manifest.tsv", ["file", "size_bytes", "mtime_epoch"], manifest_rows)


if __name__ == "__main__":
    main()
