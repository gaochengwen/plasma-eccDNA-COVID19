#!/usr/bin/env python3
"""Add Circle-Map robustness supplementary tables to the revision workbook."""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path
from typing import Dict, Iterable, List, Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REVISE = Path(__file__).resolve().parent.parent
ROOT = REVISE.parent / "analyse" / "Circle_finder"
TABLES = ROOT / "tables"
WORKBOOK = REVISE / "Supplementary_Tables_revised.xlsx"
BACKUP = REVISE / "Supplementary_Tables_revised.pre_circlemap_bak.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="FCE4D6")

CALLSET_LABELS = {
    "circlemap_current": "Current Circle-Map",
    "circlemap_methods": "Methods-consistent Circle-Map",
    "circlemap_strict": "Circle-Map split>=10, score>1,000",
    "circlemap_very_strict": "Circle-Map split>=20, score>2,000",
    "circlemap_artifact_masked": "Circle-Map artifact-masked",
    "circlemap_high_support_masked": "Circle-Map high-support + artifact-masked",
    "consensus_t10": "Circle-Map/Circle_finder consensus (10 bp)",
}


def read_tsv(name: str) -> List[Dict[str, str]]:
    path = TABLES / name
    if not path.exists():
        raise FileNotFoundError(path)
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def read_json_rows(path: Path) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows: List[Dict[str, str]] = []

    def walk(prefix: str, value) -> None:
        if isinstance(value, dict):
            for key, subvalue in value.items():
                walk(f"{prefix}.{key}" if prefix else str(key), subvalue)
        elif isinstance(value, list):
            rows.append({"parameter": prefix, "value": "; ".join(map(str, value))})
        else:
            rows.append({"parameter": prefix, "value": str(value)})

    walk("", payload)
    return rows


def as_number(value: str):
    if value is None or value == "":
        return ""
    if not isinstance(value, str):
        return value
    try:
        number = float(value)
    except ValueError:
        return value
    if number.is_integer() and "." not in value and "e" not in value.lower() and abs(number) < 1e15:
        return int(number)
    return number


def write_block(
    sheet,
    row: int,
    title: str,
    records: Sequence[Dict[str, str]],
    columns: Sequence[str] | None = None,
    note: str = "",
) -> int:
    title_cell = sheet.cell(row=row, column=1, value=title)
    title_cell.font = Font(bold=True, size=11)
    title_cell.fill = SECTION_FILL
    title_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    row += 1

    if note:
        note_cell = sheet.cell(row=row, column=1, value=note)
        note_cell.font = Font(italic=True, size=9)
        note_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        row += 1

    if not records:
        sheet.cell(row=row, column=1, value="No records")
        return row + 3

    if columns is None:
        columns = list(records[0].keys())

    for col_idx, name in enumerate(columns, 1):
        cell = sheet.cell(row=row, column=col_idx, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1

    for record in records:
        for col_idx, name in enumerate(columns, 1):
            sheet.cell(row=row, column=col_idx, value=as_number(record.get(name, "")))
        row += 1
    return row + 2


def finish_sheet(sheet) -> None:
    max_col = max(sheet.max_column, 1)
    for col_idx in range(1, max_col + 1):
        width = 24 if col_idx <= 2 else 17
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left" if cell.column <= 2 else "center",
                vertical="top",
                wrap_text=True,
            )
            if isinstance(cell.value, float):
                cell.number_format = "0.000000E+00" if cell.value and abs(cell.value) < 0.001 else "0.0000"
    sheet.freeze_panes = "A5"


def add_title(sheet, title: str) -> int:
    cell = sheet.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=12)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    return 3


def callset_definition_rows() -> List[Dict[str, str]]:
    return [
        {
            "callset": "circlemap_current",
            "label": CALLSET_LABELS["circlemap_current"],
            "primary_caller": "Circle-Map",
            "split_reads_min": "not additionally filtered in this robustness table",
            "score_filter": "not additionally filtered in this robustness table",
            "coverage_filters": "current analysis output",
            "artifact_mask": "No",
            "consensus_rule": "No",
            "notes": "Original current Circle-Map output used for comparison with the manuscript methods filter.",
        },
        {
            "callset": "circlemap_methods",
            "label": CALLSET_LABELS["circlemap_methods"],
            "primary_caller": "Circle-Map",
            "split_reads_min": ">=3",
            "score_filter": ">200",
            "coverage_filters": "mean coverage > SD; start/end coverage increase >=0.3; uncovered fraction <0.1",
            "artifact_mask": "No",
            "consensus_rule": "No",
            "notes": "Matches the high-confidence filter described in Methods.",
        },
        {
            "callset": "circlemap_strict",
            "label": CALLSET_LABELS["circlemap_strict"],
            "primary_caller": "Circle-Map",
            "split_reads_min": ">=10",
            "score_filter": ">1,000",
            "coverage_filters": "Methods coverage filters retained",
            "artifact_mask": "No",
            "consensus_rule": "No",
            "notes": "Split-read/score threshold sensitivity analysis.",
        },
        {
            "callset": "circlemap_very_strict",
            "label": CALLSET_LABELS["circlemap_very_strict"],
            "primary_caller": "Circle-Map",
            "split_reads_min": ">=20",
            "score_filter": ">2,000",
            "coverage_filters": "Methods coverage filters retained",
            "artifact_mask": "No",
            "consensus_rule": "No",
            "notes": "More stringent split-read/score threshold sensitivity analysis.",
        },
        {
            "callset": "circlemap_artifact_masked",
            "label": CALLSET_LABELS["circlemap_artifact_masked"],
            "primary_caller": "Circle-Map",
            "split_reads_min": ">=3",
            "score_filter": ">200",
            "coverage_filters": "Methods coverage filters retained",
            "artifact_mask": "No full-circle or +/-100 bp breakpoint overlap with blacklist/gap; breakpoint Umap k100 fraction >=0.90; no segmental duplication or RepeatMasker overlap",
            "consensus_rule": "No",
            "notes": "Internal repeats were retained; only breakpoint and full-circle artifact masks were applied.",
        },
        {
            "callset": "circlemap_high_support_masked",
            "label": CALLSET_LABELS["circlemap_high_support_masked"],
            "primary_caller": "Circle-Map",
            "split_reads_min": ">=10",
            "score_filter": ">1,000",
            "coverage_filters": "Methods coverage filters retained",
            "artifact_mask": "Same as circlemap_artifact_masked",
            "consensus_rule": "No",
            "notes": "High-support call set used as the strict single-caller sensitivity set.",
        },
        {
            "callset": "consensus_t10",
            "label": CALLSET_LABELS["consensus_t10"],
            "primary_caller": "Circle-Map and Circle_finder",
            "split_reads_min": "Circle-Map methods filter; Circle_finder support >=2",
            "score_filter": "Circle-Map score >200",
            "coverage_filters": "Circle-Map methods coverage filters retained",
            "artifact_mask": "No additional artifact mask unless specified in target audit",
            "consensus_rule": "same chromosome and both breakpoints within 10 bp; one-to-one matching",
            "notes": "Primary two-caller consensus call set used for downstream robustness analyses.",
        },
    ]


def build_s12(workbook) -> None:
    sheet = workbook.create_sheet("Table_S12")
    row = add_title(
        sheet,
        "Supplementary Table S12. Circle-Map call-set sensitivity analysis for COVID-19 versus HC eccDNA burden.",
    )
    row = write_block(sheet, row, "S12a. Call-set definitions", callset_definition_rows())
    row = write_block(
        sheet,
        row,
        "S12b. Per-sample burden metrics across call sets",
        read_tsv("robustness_callset_sample_metrics.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S12c. Unadjusted group tests for call count, EPM and retention fraction",
        read_tsv("robustness_group_tests.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S12d. Robust covariate-adjusted regression models with HC3 standard errors",
        read_tsv("robustness_regression_HC3.tsv"),
        note="The group term estimates COVID-19 versus HC effects in log2(EPM + 1), adjusted for log2 RCA product concentration, age and sex.",
    )
    row = write_block(
        sheet,
        row,
        "S12e. Circle-Map split-read and score threshold grid",
        read_tsv("robustness_threshold_grid_tests.tsv"),
    )
    finish_sheet(sheet)


def build_s13(workbook) -> None:
    sheet = workbook.create_sheet("Table_S13")
    row = add_title(
        sheet,
        "Supplementary Table S13. Circle_finder validation, two-caller concordance and consensus provenance.",
    )
    row = write_block(
        sheet,
        row,
        "S13a. Circle_finder per-sample validation",
        read_tsv("circlefinder_validation_all.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S13b. Per-sample Circle-Map/Circle_finder concordance across breakpoint tolerances",
        read_tsv("caller_concordance_all.tsv"),
        note="Consensus requires the same chromosome and both breakpoints within the stated tolerance, matched one-to-one.",
    )
    row = write_block(
        sheet,
        row,
        "S13c. Circle_finder and consensus software provenance",
        read_tsv("software_versions.tsv"),
    )
    parameter_rows = []
    for name in [
        "circlefinder_validation_parameters_all.json",
        "consensus_run_parameters_all.json",
        "core_robustness_run_parameters.json",
    ]:
        for record in read_json_rows(ROOT / name):
            parameter_rows.append({"source_file": name, **record})
    row = write_block(sheet, row, "S13d. Key run parameters serialized from analysis JSON files", parameter_rows)
    finish_sheet(sheet)


def build_s14(workbook) -> None:
    sheet = workbook.create_sheet("Table_S14")
    row = add_title(
        sheet,
        "Supplementary Table S14. Fragment-size sensitivity analyses across robust Circle-Map and consensus call sets.",
    )
    row = write_block(
        sheet,
        row,
        "S14a. Group tests for length-bin proportions",
        read_tsv("robustness_fragment_size_group_tests.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S14b. Per-sample length-bin proportions",
        read_tsv("robustness_fragment_size_sample_proportions.tsv"),
    )
    finish_sheet(sheet)


def build_s15(workbook) -> None:
    sheet = workbook.create_sheet("Table_S15")
    row = add_title(
        sheet,
        "Supplementary Table S15. Downstream eccGene and chromatin robustness under high-support filtering and two-caller consensus calling.",
    )
    row = write_block(sheet, row, "S15a. eccGene tested and significant-gene counts", read_tsv("eccgene_callset_summary.tsv"))
    row = write_block(
        sheet,
        row,
        "S15b. eccGene concordance relative to methods-consistent Circle-Map",
        read_tsv("eccgene_callset_concordance.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S15c. Candidate eccGene stability for CTNNA2, CAB39 and SDK1",
        read_tsv("eccgene_candidate_stability.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S15d. Chromatin effect concordance relative to methods-consistent Circle-Map",
        read_tsv("chromatin_callset_concordance.tsv"),
        note="Each row compares the COVID-19 minus HC median difference for a matched peak set, localization mode and metric.",
    )
    row = write_block(
        sheet,
        row,
        "S15e. Chromatin group statistics recalculated under each robustness call set",
        read_tsv("chromatin_callset_group_statistics.tsv"),
    )
    finish_sheet(sheet)


def build_s16(workbook) -> None:
    sheet = workbook.create_sheet("Table_S16")
    row = add_title(
        sheet,
        "Supplementary Table S16. Validation-target support and strict breakpoint artifact-mask audit.",
    )
    row = write_block(
        sheet,
        row,
        "S16a. Target support summary by call set, tolerance and cohort",
        read_tsv("validation_target_support_summary.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S16b. Per-sample target support records",
        read_tsv("validation_target_sample_support.tsv"),
    )
    row = write_block(
        sheet,
        row,
        "S16c. Strict breakpoint artifact-mask audit for each target breakpoint",
        read_tsv("validation_target_mask_audit.tsv"),
        note="Both 200-bp breakpoint windows must have Umap k100 unique fraction >=0.90 and no blacklist/gap, segmental-duplication or RepeatMasker overlap to pass the strict target-level mask.",
    )
    row = write_block(
        sheet,
        row,
        "S16d. Artifact-filter stage counts by sample",
        read_tsv("artifact_filter_stage_counts.tsv"),
    )
    finish_sheet(sheet)


def rebuild_circlemap_sheets() -> None:
    if not BACKUP.exists():
        shutil.copy2(WORKBOOK, BACKUP)
    workbook = load_workbook(WORKBOOK)
    for sheet_name in [f"Table_S{i}" for i in range(12, 17)]:
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
    build_s12(workbook)
    build_s13(workbook)
    build_s14(workbook)
    build_s15(workbook)
    build_s16(workbook)
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    rebuild_circlemap_sheets()
