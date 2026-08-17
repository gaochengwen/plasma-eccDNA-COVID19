#!/usr/bin/env python3
"""Package the audited TSV outputs into a styled supplementary workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = [
    ("Selected_Loci", "selected_loci.tsv"),
    ("BCL3_PROCR_Audit", "BCL3_PROCR_audit.tsv"),
    ("Top10_By_Category", "top10_by_category.tsv"),
    ("Category_Rankings", "category_candidate_rankings.tsv"),
    ("All_Candidates", "candidate_locus_ranking_all.tsv"),
    ("Ranking_Criteria", "ranking_criteria.tsv"),
    ("Category_Definitions", "category_definitions.tsv"),
    ("Analysis_Summary", "current_vs_quality_filtered_summary.tsv"),
    ("Sample_Abundance", "sample_abundance_selected_and_audit.tsv"),
]


def style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="3C5488")
    header_font = Font(name="Arial", size=8, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=8, color="000000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 36
    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for index, column in enumerate(worksheet.iter_cols(), start=1):
        header = str(column[0].value or "")
        observed = len(header)
        for cell in column[1 : min(len(column), 201)]:
            observed = max(observed, len(str(cell.value or "")))
        if any(token in header.lower() for token in ("reason", "rule", "evidence", "scope")):
            width = min(max(observed + 2, 24), 60)
        elif header.lower() in {"gene", "functional_category", "category_for_row"}:
            width = min(max(observed + 2, 14), 30)
        else:
            width = min(max(observed + 2, 11), 24)
        worksheet.column_dimensions[get_column_letter(index)].width = width


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tables-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        for sheet_name, filename in SHEETS:
            frame = pd.read_csv(args.tables_dir / filename, sep="\t")
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            style_sheet(writer.book[sheet_name])
        properties = writer.book.properties
        properties.title = "Objective eccGene locus-selection ranking"
        properties.subject = "BCL3/PROCR selection-bias audit"
        properties.creator = "COVID-19 eccDNA revision analysis"
        properties.description = (
            "Lexicographic candidate ranking; chromatin signal and IGV appearance "
            "were not used for selection."
        )
    print(f"Wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
