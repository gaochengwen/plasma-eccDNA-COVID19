#!/usr/bin/env python3
"""Check every chromatin number in the manuscript and the three response letters
against Supplementary Tables S9-S11 and the rendered figures.

Each assertion recomputes the value from the delivered workbook or reads it out
of the figure's SVG rather than trusting the prose, so the check fails if the
analysis is rerun and the text is not updated with it. Run after any edit to the
chromatin section.
"""

from __future__ import annotations

import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_tracked_edit import accepted_text, load_document  # noqa: E402

REVISE = Path(__file__).resolve().parent.parent
WORKBOOK = REVISE / "Supplementary_Tables_revised.xlsx"
DOC = REVISE / "Maintext_Covid-eccDNA_revised_tracked_v3.docx"
RESPONSES = {
    "R1": REVISE / "JARE-D-26-04883_Reviewer_1_response.docx",
    "R2": REVISE / "JARE-D-26-04883_Reviewer_2_response.docx",
    "R3": REVISE / "JARE-D-26-04883_Reviewer_3_response.docx",
}
SVG = "{http://www.w3.org/2000/svg}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

CD14_MARKS = ("H3K27ac", "H3K27me3", "H3K4me1", "H3K4me3", "H3K9ac", "H3K9me3")

results: List[Tuple[bool, str]] = []


def check(ok: bool, message: str) -> None:
    results.append((bool(ok), message))


HEADER_KEYS = ("peak_set_id", "sample_id")


def section(ws, label: str) -> Tuple[int, List[str]]:
    """Header row index and column names of a labelled block such as 'S9d.'.

    Each block is a label row, an optional note row, then the header row, which
    is identified by its first cell rather than by position because the number
    of note rows varies between blocks.
    """
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=1).value
        if not (isinstance(value, str) and value.startswith(label)):
            continue
        for header in range(row + 1, min(row + 8, ws.max_row + 1)):
            if ws.cell(row=header, column=1).value in HEADER_KEYS:
                cells = [ws.cell(row=header, column=c).value for c in range(1, ws.max_column + 1)]
                return header, [str(c) if c is not None else "" for c in cells]
        raise KeyError(f"no header row under {label!r} in {ws.title}")
    raise KeyError(f"block {label!r} not found in {ws.title}")


def rows_of(ws, label: str) -> List[Dict[str, object]]:
    header_row, columns = section(ws, label)
    out: List[Dict[str, object]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        first = ws.cell(row=row, column=1).value
        if first is None:
            break
        if isinstance(first, str) and re.match(r"^S\d+[a-z]\.", first):
            break
        record = {
            columns[c - 1]: ws.cell(row=row, column=c).value
            for c in range(1, len(columns) + 1)
        }
        out.append(record)
    return out


def benjamini_hochberg(pvalues: List[float]) -> List[float]:
    n = len(pvalues)
    order = sorted(range(n), key=lambda i: pvalues[i])
    out = [0.0] * n
    running = 1.0
    for rank, index in enumerate(reversed(order), start=1):
        running = min(running, pvalues[index] * n / (n - rank + 1))
        out[index] = running
    return out


def bh_reproduces(ws, label: str, p_col: str, q_col: str, family_cols) -> bool:
    """Recompute BH inside the family the Methods claim and compare with the table."""
    groups: Dict[tuple, List[dict]] = {}
    for record in rows_of(ws, label):
        key = tuple(str(record[c]) for c in family_cols)
        groups.setdefault(key, []).append(record)
    for members in groups.values():
        expected = benjamini_hochberg([float(r[p_col]) for r in members])
        for record, value in zip(members, expected):
            if abs(float(record[q_col]) - value) > max(1e-12, abs(value) * 1e-9):
                return False
    return True


def close(a, b, tol) -> bool:
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def svg_texts(path: Path) -> List[str]:
    root = ET.parse(path).getroot()
    return ["".join(node.itertext()).strip() for node in root.iter(SVG + "text")]


def docx_text(path: Path) -> str:
    root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    return "".join(node.text or "" for node in root.iter(W + "t"))


def main() -> int:
    wb = load_workbook(WORKBOOK, data_only=True)
    s9, s10, s11 = wb["Table_S9"], wb["Table_S10"], wb["Table_S11"]
    text = accepted_text(load_document(DOC))
    letters = {name: docx_text(path) for name, path in RESPONSES.items()}
    everything = text + "\n" + "\n".join(letters.values())

    # ---- S9c: within-cohort enrichment ---------------------------------
    within = [r for r in rows_of(s9, "S9c.")
              if r["analysis_tier"] == "primary" and r["mode"] == "full_interval"]
    check(len(within) == 12, f"S9c has 12 primary full-interval rows (found {len(within)})")
    ratios = {(r["mark"], r["group"]): float(r["median_enrichment_ratio"]) for r in within}
    enriched = [v for (mark, _), v in ratios.items() if mark != "H3K9me3"]
    lo, hi = min(enriched), max(enriched)
    check(f"{lo:.2f}–{hi:.2f}" == "1.16–1.30",
          f"stated enrichment range 1.16-1.30 matches the tables ({lo:.3f}-{hi:.3f})")
    check("Median enrichment ratios were 1.16–1.30" in text,
          "manuscript quotes the 1.16-1.30 enrichment range")
    check("median enrichment ratios of 1.16–1.30" in letters["R2"],
          "Reviewer 2 letter quotes the 1.16-1.30 enrichment range")
    check(all(float(r["bh_fdr_within_family"]) < 4e-12 for r in within),
          "all twelve within-cohort tests are at q = 3.64e-12")
    check(close(ratios[("H3K9me3", "COVID")], 0.766, 0.001)
          and close(ratios[("H3K9me3", "HC")], 0.804, 0.001),
          "H3K9me3 ratios quoted as 0.77 (COVID-19) and 0.80 (HC)")

    # ---- S9d/S9e: between-cohort contrasts ------------------------------
    rel = {r["mark"]: r for r in rows_of(s9, "S9d.")
           if r["analysis_tier"] == "primary" and r["mode"] == "full_interval"}
    check(close(rel["H3K27me3"]["cliffs_delta_covid_vs_hc"], 0.675, 0.001)
          and close(rel["H3K27me3"]["bh_fdr_all_tests"], 8.9e-6, 5e-8),
          "H3K27me3 between-cohort delta 0.675 / q 8.9e-6")
    check(close(rel["H3K9me3"]["cliffs_delta_covid_vs_hc"], -0.552, 0.001)
          and close(rel["H3K9me3"]["bh_fdr_all_tests"], 1.7e-4, 5e-6),
          "H3K9me3 between-cohort delta -0.552 / q 1.7e-4")
    check(close(rel["H3K9ac"]["cliffs_delta_covid_vs_hc"], 0.324, 0.001)
          and close(rel["H3K9ac"]["bh_fdr_all_tests"], 0.037, 0.001),
          "H3K9ac between-cohort delta 0.324 / q 0.037")
    # H3K9ac is significant only under full-interval, which the text must say.
    h3k9ac_modes = {r["mode"]: float(r["bh_fdr_all_tests"]) for r in rows_of(s9, "S9d.")
                    if r["mark"] == "H3K9ac" and r["analysis_tier"] == "primary"}
    check(h3k9ac_modes["full_interval"] < 0.05
          and min(h3k9ac_modes["junction_start_end"], h3k9ac_modes["midpoint"]) > 0.05,
          "H3K9ac reaches q<0.05 only under the full-interval definition")
    check("the weaker H3K9ac difference reached significance only under the full-interval "
          "definition" in text,
          "Methods state the H3K9ac definition dependence")
    check("weaker H3K9ac difference reaches significance only under the full-interval definition"
          in letters["R3"],
          "Reviewer 3 letter states the H3K9ac definition dependence")

    epm = {r["mark"]: r for r in rows_of(s9, "S9e.")
           if r["analysis_tier"] == "primary" and r["mode"] == "full_interval"}
    deltas = [float(epm[m]["cliffs_delta_covid_vs_hc"]) for m in CD14_MARKS]
    qs = [float(epm[m]["bh_fdr_all_tests"]) for m in CD14_MARKS]
    check(f"{min(deltas):.3f} to {max(deltas):.3f}" == "0.740 to 0.769",
          f"absolute-abundance delta range 0.740-0.769 ({min(deltas):.4f}-{max(deltas):.4f})")
    check(min(qs) >= 2.2e-8 and max(qs) <= 2.95e-8,
          f"absolute-abundance q range 2.2e-8 to 2.9e-8 ({min(qs):.3g}-{max(qs):.3g})")

    a549 = {(r["mark"], r["mode"]): r for r in rows_of(s9, "S9e.")
            if r["analysis_tier"] == "secondary"}
    check(close(a549[("H3K27ac", "full_interval")]["median_covid"], 71.0, 0.05)
          and close(a549[("H3K27ac", "full_interval")]["median_hc"], 17.9, 0.05),
          "A549 H3K27ac medians 71.0 vs 17.9")
    check(close(a549[("H3K4me3", "full_interval")]["median_covid"], 43.7, 0.05)
          and close(a549[("H3K4me3", "full_interval")]["median_hc"], 12.1, 0.05),
          "A549 H3K4me3 medians 43.7 vs 12.1")
    a549_rel = {(r["mark"], r["mode"]): r for r in rows_of(s9, "S9d.")
                if r["analysis_tier"] == "secondary"}
    check(close(a549_rel[("H3K4me3", "junction_start_end")]["cliffs_delta_covid_vs_hc"],
                -0.335, 0.001)
          and close(a549_rel[("H3K4me3", "midpoint")]["cliffs_delta_covid_vs_hc"], -0.324, 0.001),
          "A549 H3K4me3 junction/midpoint deltas -0.335 and -0.324")
    check("reached significance only under the junction and midpoint definitions" in text,
          "Results describe the A549 H3K4me3 definition dependence honestly")

    # ---- S10: burden controls -------------------------------------------
    down = {(r["mark"], r["mode"], r["downsample_target"]): r for r in rows_of(s10, "S10a.")
            if r["analysis_tier"] == "primary"}
    k27 = down[("H3K27me3", "full_interval", 4000)]
    k9 = down[("H3K9me3", "full_interval", 4000)]
    check(close(k27["bh_fdr_within_family"], 1.9e-6, 5e-8),
          f"downsampled H3K27me3 q is 1.9e-6 ({float(k27['bh_fdr_within_family']):.3g})")
    check(close(k9["bh_fdr_within_family"], 6.5e-5, 5e-7),
          f"downsampled H3K9me3 q is 6.5e-5 ({float(k9['bh_fdr_within_family']):.3g})")
    check("q = 1.9 × 10⁻⁶" in text and "q = 6.5 × 10⁻⁵" in text,
          "manuscript quotes the current downsampled q values")
    check("q = 1.9 × 10⁻⁶" in letters["R2"] and "q = 6.5 × 10⁻⁵" in letters["R2"],
          "Reviewer 2 letter quotes the current downsampled q values")
    check("1.5 × 10⁻⁵" not in everything, "the superseded 1.5e-5 q value appears nowhere")

    matched = [r for r in rows_of(s10, "S10c.")
               if r["analysis_tier"] == "primary" and r["mode"] == "full_interval"]
    check(min(float(r["bh_fdr_within_family"]) for r in matched) >= 0.55,
          "all burden-matched comparisons have q >= 0.55")
    check(all(int(r["n_covid"]) == 11 and int(r["n_hc"]) == 18 for r in matched)
          and close(matched[0]["residual_burden_difference_p"], 0.84, 0.005),
          "burden-matched window is 11 vs 18 with residual P = 0.84")

    hc3 = rows_of(s10, "S10e.")
    m3_group = [r for r in hc3 if r["model"] == "m3_plus_eccdna_burden"
                and r["term"] == "COVID_vs_HC" and r["mode"] == "full_interval"
                and r["analysis_tier"] == "primary"]
    check(len(m3_group) == 6 and min(float(r["bh_fdr_within_family"]) for r in m3_group) >= 0.97,
          "group term q >= 0.97 for all six marks once burden is in the model")
    check(all(close(r["vif"], 4.09, 0.01) for r in m3_group),
          "group-term VIF is 4.09 in the burden-adjusted model")

    # ---- S11: cross-cell-type replication --------------------------------
    provenance = rows_of(s11, "S11a.")
    check(len(provenance) == 33, f"S11a lists 33 GRCh38-native peak sets (found {len(provenance)})")
    check("33 histone ChIP-seq peak sets released by ENCODE natively on GRCh38" in text,
          "manuscript quotes 33 GRCh38-native peak sets")

    cross = [r for r in rows_of(s11, "S11b.") if r["mode"] == "full_interval"]
    for mark, expected in (("H3K27ac", "1.24–1.38"), ("H3K4me1", "1.25–1.31")):
        vals = [float(r["median_ratio_covid"]) for r in cross if r["mark"] == mark]
        check(f"{min(vals):.2f}–{max(vals):.2f}" == expected,
              f"{mark} cross-cell range {expected} ({min(vals):.3f}-{max(vals):.3f})")
        above = [int(r["n_above_expected_covid"]) for r in cross if r["mark"] == mark]
        check(min(above) >= 38, f"{mark}: 38 or 39 of 39 COVID-19 samples above expectation")
    hc_above = [int(r["n_above_expected_hc"]) for r in cross
                if r["mark"] in ("H3K27ac", "H3K4me1")]
    check(min(hc_above) == 36,
          f"HC counts above expectation bottom out at 36, as the letters now say ({min(hc_above)})")
    check("36 to 39 of the 39 samples in each cohort" in letters["R3"],
          "Reviewer 3 letter quotes 36 to 39 samples above expectation")

    # Peak-set breadth: the lifted files are two- to eightfold broader.
    lifted = {r["mark"]: float(r["peak_bp"]) for r in rows_of(s11, "S11c.")
              if r["analysis_tier"] == "primary"}
    native = {r["mark"]: float(r["merged_coverage_bp"]) for r in provenance
              if r["biosample"] == "CD14-positive monocyte"}
    folds = [lifted[m] / native[m] for m in CD14_MARKS]
    check(2.0 <= min(folds) < 3.0 and 7.5 < max(folds) <= 8.5,
          f"lifted peak sets are two- to eightfold broader ({min(folds):.2f}-{max(folds):.2f})")
    check("two- to eightfold less of the genome" in text,
          "manuscript states the two- to eightfold breadth ratio")
    check("between two and eight times less of the genome" in letters["R3"],
          "Reviewer 3 letter states the two- to eightfold breadth ratio")
    check("three- to eightfold" not in everything and "three and eight times" not in everything,
          "the superseded three- to eightfold claim appears nowhere")

    mapp = {r["mark"]: r for r in rows_of(s11, "S11c.") if r["analysis_tier"] == "primary"}
    check(close(mapp["H3K9me3"]["mappable_fraction"], 0.96, 0.005)
          and close(mapp["H3K9me3"]["genome_background_mappable_fraction"], 0.94, 0.005),
          "H3K9me3 mappability 0.96 versus a 0.94 background")
    restricted = {(r["mark"], r["mode"]): r for r in rows_of(s11, "S11d.")}
    check(close(restricted[("H3K9me3", "full_interval")]["median_ratio_covid"], 0.756, 0.001),
          "H3K9me3 ratio is 0.756 after restriction to mappable sequence")

    binning = [float(r["max_absolute_difference"]) for r in rows_of(s11, "S11e.")
               if str(r["peak_set_id"]).startswith("CD14")]
    check(f"{max(binning):.4f}" == "0.0063",
          f"25 bp binning changes log2 enrichment by at most 0.0063 ({max(binning):.5f})")
    check("at most 0.0063" in text, "manuscript quotes the 0.0063 binning bound")

    # ---- figure/legend agreement ----------------------------------------
    fig4 = svg_texts(REVISE / "Figure_4.svg")
    tiers4 = [t for t in fig4 if set(t) == {"*"}]
    check(len(tiers4) == 9,
          f"Figure 4 carries nine significance tiers, six in b and three in c (found {len(tiers4)})")
    check("Asterisks in (b) and (c) are two-sided Wilcoxon rank-sum tests between cohorts" in text,
          "Figure 4 legend describes the between-cohort asterisks that panel c actually shows")
    check("contrast in relative enrichment is not marked here" not in text,
          "the retracted 'not marked here' sentence is gone from the Figure 4 legend")

    fig5 = svg_texts(REVISE / "Figure_5.svg")
    labels5 = sorted({t for t in fig5 if len(t) == 1 and t.isalpha()})
    check(labels5 == ["a", "b", "c", "d", "e"],
          f"Figure 5 has panels a-e (found {labels5})")
    for phrase in ("(b, c) Absolute abundance within the H3K27ac (b) and H3K4me3 (c) peak sets",
                   "(d, e) The BCL3 (chr19:44,747,207–44,761,544) and PROCR "
                   "(chr20:35,170,571–35,178,862) loci"):
        check(phrase in text, f"Figure 5 legend describes {phrase[:34]!r}")
    check("(Fig. 5d)" in text and "(Fig. 5e)" in text and "Fig. 5b,c;" in text,
          "in-text Figure 5 callouts use the current panel letters")
    for locus in ("chr19:44,747,207", "chr20:35,170,571"):
        check(any(locus in t for t in fig5) and locus in text,
              f"legend and figure agree on {locus}")

    # Cohort colours: Figures 4, 5 and S4 all use the same red/blue pair.
    for name in ("Figure_4.svg", "Figure_5.svg", "Figure_S4.svg"):
        raw = (REVISE / name).read_text()
        check("#c84f50" in raw.lower() and "#36617b" in raw.lower(),
              f"{name} uses the shared COVID-19 red and HC blue")
    check("red denotes COVID-19 and blue HC" in text,
          "figure legends describe the cohort colours as red and blue")
    check("orange COVID-19" not in text, "no legend still calls the COVID-19 colour orange")

    # Every drawn panel of Figures 4 and 5 must be cited somewhere in the text.
    body = text[: text.index("Material and methods")]
    for panel in ("4a", "4b", "4c", "5a", "5b", "5c", "5d", "5e"):
        figure, letter = panel[0], panel[1]
        cited = re.search(rf"Fig\. ?{figure}[a-e]*,?{letter}\b", body) is not None
        check(cited, f"Results cite Fig. {panel}")

    # Legends stay within a length a journal will accept without a fight.
    for name, start, end, limit in (
        ("Figure 4", "Figure 4. Plasma eccDNAs associate", "**** q < 0.0001.", 1700),
        ("Figure 5", "Figure 5. Plasma eccDNA abundance", "**** q < 0.0001.", 1700),
        ("Figure S4", "Supplementary Figure 4. Relative enrichment at CD14+",
         "Supplementary Tables S9 and S10.", 2200),
        ("Figure S5", "Supplementary Figure 5. The chromatin association",
         "Supplementary Table S11.", 1800),
    ):
        i = text.index(start)
        length = text.index(end, i) + len(end) - i
        check(length <= limit, f"{name} legend is {length} characters (limit {limit})")

    # ---- supplementary material list -------------------------------------
    titles = {sheet.title: str(sheet.cell(row=1, column=1).value) for sheet in wb.worksheets}
    check(len(titles) == 11 and "Table_S11" in titles, "workbook holds a contiguous S1-S11")
    for number, keyword in ((8, "Two-sided COVID-19 versus HC comparisons"),
                            (9, "Chromatin analysis in a single hg38 coordinate system"),
                            (10, "Controls for the difference in detected eccDNA burden"),
                            (11, "Cross-cell-type replication using GRCh38-native ENCODE histone")):
        check(f"Supplementary Table {number}. {keyword}" in text,
              f"supplementary list entry S{number} matches the workbook title")
        check(keyword in titles[f"Table_S{number}"],
              f"workbook sheet Table_S{number} carries the same title")
    for legend in ("Supplementary Figure 4. Relative enrichment at CD14+ histone-mark peak sets",
                   "Supplementary Figure 5. The chromatin association reproduces"):
        check(legend in text, f"manuscript carries the {legend[:24]} legend")

    # ---- every reported analysis is described in Methods ------------------
    # Each entry is (what the Results or a legend reports, the Methods phrase
    # that has to describe it).
    for reported, described in (
        ("one-sample Wilcoxon signed-rank test",
         "one-sample two-sided Wilcoxon signed-rank test of the per-sample log2 enrichment "
         "against 0"),
        ("percentile bootstrap 95% confidence interval of the median",
         "percentile bootstrap intervals over 10,000 resamples"),
        ("released with GSE179184",
         "From GSE179184 we used the processed H3K27ac and H3K4me3 peak sets"),
        ("in-house thresholded A549 peak files",
         "labelled legacy sensitivity tier that supports no conclusion"),
        ("33 histone ChIP-seq peak sets released by ENCODE natively on GRCh38",
         "released natively on GRCh38 by ENCODE"),
        ("uniquely mappable", "the Umap k36 track (Bismap)"),
        ("25 bp bins", "rounded into 25 bp bins"),
        ("downsampling every sample", "randomly reduced without replacement to a common number"),
        ("burden-matched count window", "within the range of eligible eccDNA counts shared by "
                                        "both cohorts"),
        ("HC3 robust regression", "ordinary least squares with HC3 robust standard errors"),
        ("Spearman", "quantified by Spearman correlation within each cohort separately"),
    ):
        check(described in text, f"Methods describe {reported!r}")

    # Table S8 lost its histone rows with the superseded metric, so nothing may
    # still describe it as holding histone-mark families.
    families = {
        s8.cell(row=r, column=1).value
        for r in range(3, s8.max_row + 1)
        if s8.cell(row=r, column=1).value
    } if (s8 := wb["Table_S8"]) else set()
    check(not any("histone" in str(f) for f in families),
          f"Table S8 no longer contains histone-mark families ({sorted(families)})")
    check("histone-mark metric families" not in text,
          "Statistical analysis no longer lists histone families inside Table S8")
    check("The chromatin analyses form their own families" in text,
          "Statistical analysis documents the chromatin BH families")

    # The stated BH families must reproduce the stored q values exactly.
    check(bh_reproduces(s9, "S9c.", "wilcoxon_signed_rank_p_vs_0", "bh_fdr_within_family",
                        ("analysis_tier", "mode", "group")),
          "within-cohort q values reproduce under the family stated in Methods")
    check(bh_reproduces(s9, "S9d.", "mannwhitneyu_p", "bh_fdr_all_tests", ()),
          "relative-enrichment q values reproduce as one 30-test family")
    check(bh_reproduces(s10, "S10e.", "p_value", "bh_fdr_within_family",
                        ("analysis_tier", "mode", "model", "term")),
          "HC3 regression q values reproduce under the family stated in Methods")
    check(bh_reproduces(s11, "S11b.", "mannwhitneyu_p", "bh_fdr_group_within_family", ("mode",)),
          "cross-cell-type q values reproduce under the family stated in Methods")

    # ---- claims that must no longer appear anywhere -----------------------
    for retracted in ("preferentially localize to active regulatory chromatin",
                      "reduced relative enrichment across multiple chromatin features",
                      "relative enrichment was significantly reduced for only three of the four"):
        check(retracted not in everything, f"retracted claim absent: {retracted[:44]!r}")

    passed = sum(1 for ok, _ in results if ok)
    for ok, message in results:
        print(("PASS: " if ok else "FAIL: ") + message)
    print(f"\n{passed}/{len(results)} checks passed")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
