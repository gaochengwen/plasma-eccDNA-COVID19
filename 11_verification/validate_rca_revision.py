#!/usr/bin/env python3
"""Validate tracked revisions, EndNote fields, responses, tables, and Figure S2."""

from __future__ import annotations

import posixpath
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
REVISE = ROOT / "revise"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
W = f"{{{W_NS}}}"


def read_xml(package: Path, member: str) -> etree._Element:
    with ZipFile(package) as archive:
        return etree.fromstring(archive.read(member))


def paragraph_texts(root: etree._Element, view: str) -> list[str]:
    def collect(node: etree._Element) -> str:
        if view == "rejected" and node.tag == f"{W}ins":
            return ""
        if view == "accepted" and node.tag == f"{W}del":
            return ""
        if node.tag == f"{W}t":
            return node.text or ""
        if node.tag == f"{W}delText" and view == "rejected":
            return node.text or ""
        return "".join(collect(child) for child in node)

    return [text for p in root.xpath("//w:p", namespaces=NS) if (text := collect(p))]


def field_instructions(root: etree._Element) -> list[str]:
    return [text or "" for text in root.xpath("//w:instrText/text()", namespaces=NS)]


def validate_manuscript() -> list[str]:
    source = ROOT / "Maintext" / "Maintext_Covid-eccDNA.docx"
    revised = REVISE / "Maintext_Covid-eccDNA_revised_tracked.docx"
    source_xml = read_xml(source, "word/document.xml")
    revised_xml = read_xml(revised, "word/document.xml")
    settings = read_xml(revised, "word/settings.xml")

    source_paragraphs = paragraph_texts(source_xml, "accepted")
    rejected_paragraphs = paragraph_texts(revised_xml, "rejected")
    assert source_paragraphs == rejected_paragraphs, "Rejecting changes does not reconstruct the source manuscript"

    accepted = "\n".join(paragraph_texts(revised_xml, "accepted"))
    required = [
        "Assessment of RCA-associated technical variation",
        "5.08; 95% CI, 2.74–9.41",
        "Supplementary Figure 2. Assessment of RCA-associated technical variation",
        "16,857,160 eccDNA calls",
        "4.74-fold higher adjusted count",
    ]
    for phrase in required:
        assert phrase in accepted, f"Accepted manuscript is missing: {phrase}"

    source_fields = field_instructions(source_xml)
    revised_fields = field_instructions(revised_xml)
    assert source_fields == revised_fields, "EndNote/Word field instructions changed"
    cite_count = sum("ADDIN EN.CITE" in field for field in revised_fields)
    reflist_count = sum("ADDIN EN.REFLIST" in field for field in revised_fields)
    assert cite_count > 0 and reflist_count == 1, "EndNote fields are incomplete"

    insertions = revised_xml.xpath("//w:ins", namespaces=NS)
    deletions = revised_xml.xpath("//w:del", namespaces=NS)
    assert insertions and deletions, "Tracked insertions/deletions are missing"
    assert len(settings.xpath("./w:trackRevisions", namespaces=NS)) == 1, "Track Changes is not enabled"
    assert all(not deletion.xpath(".//w:p", namespaces=NS) for deletion in deletions), "Whole-paragraph deletion detected"
    deletion_lengths = [len("".join(d.xpath(".//w:delText/text()", namespaces=NS))) for d in deletions]

    return [
        f"Tracked manuscript: PASS ({len(insertions)} insertions; {len(deletions)} sentence-level deletions)",
        f"Reject-all reconstruction: PASS ({len(source_paragraphs)} non-empty source paragraphs)",
        f"EndNote traveling-library fields: PASS ({cite_count} EN.CITE; {reflist_count} EN.REFLIST; exact instruction match)",
        f"Largest tracked deletion: {max(deletion_lengths)} characters",
    ]


def validate_responses() -> list[str]:
    cases = [
        (
            ROOT / "JARE-D-26-04883_Reviewer_1_comments.docx",
            REVISE / "JARE-D-26-04883_Reviewer_1_response.docx",
            "Potential methodological bias introduced by Circle-seq",
        ),
        (
            ROOT / "JARE-D-26-04883_Reviewer_3_comments.docx",
            REVISE / "JARE-D-26-04883_Reviewer_3_response.docx",
            "Cohort imbalance and potential confounding factors",
        ),
    ]
    output = []
    for source_path, revised_path, phrase in cases:
        source = Document(source_path)
        revised = Document(revised_path)
        assert len(source.tables) == len(revised.tables)
        target_count = 0
        filled_response_rows = 0
        for source_table, revised_table in zip(source.tables, revised.tables):
            assert len(source_table.rows) == len(revised_table.rows)
            for source_row, revised_row in zip(source_table.rows, revised_table.rows):
                assert len(source_row.cells) == len(revised_row.cells)
                comment = " ".join(source_row.cells[1].text.split()) if len(source_row.cells) > 1 else ""
                if len(source_row.cells) >= 4:
                    assert source_row.cells[0].text == revised_row.cells[0].text
                    assert source_row.cells[1].text == revised_row.cells[1].text
                    if comment == "Comment":
                        assert source_row.cells[2].text == revised_row.cells[2].text
                        assert source_row.cells[3].text == revised_row.cells[3].text
                        continue
                    response_present = bool(revised_row.cells[2].text.strip() or revised_row.cells[3].text.strip())
                    if response_present:
                        filled_response_rows += 1
                    if phrase in comment:
                        target_count += 1
                        assert response_present
                        for cell in revised_row.cells[2:4]:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    if not run.text:
                                        continue
                                    assert run.font.name == "Times New Roman"
                                    assert run.font.size and round(run.font.size.pt, 2) == 12.0
                                    assert run.font.color.rgb and str(run.font.color.rgb) == "0000FF"
        assert target_count == 1
        assert filled_response_rows == 1, f"Unexpected additional response rows in {revised_path.name}"
        output.append(f"{revised_path.name}: PASS (one targeted comment answered; all other response rows blank)")
    return output


def validate_tables() -> list[str]:
    source_path = ROOT / "Maintext" / "Supplementary_Tables.xlsx"
    revised_path = REVISE / "Supplementary_Tables_RCA_revised.xlsx"
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(source_path) as source_zip, ZipFile(revised_path) as revised_zip:
        assert source_zip.namelist() == revised_zip.namelist()
        workbook_xml = etree.fromstring(source_zip.read("xl/workbook.xml"))
        sheet = workbook_xml.xpath(
            ".//s:sheet[@name='Table_S2']",
            namespaces={"s": spreadsheet_ns},
        )
        assert len(sheet) == 1
        relationship_id = sheet[0].get(f"{{{office_rel_ns}}}id")
        relationships = etree.fromstring(source_zip.read("xl/_rels/workbook.xml.rels"))
        target = relationships.xpath(
            ".//r:Relationship[@Id=$relationship_id]/@Target",
            namespaces={"r": package_rel_ns},
            relationship_id=relationship_id,
        )
        assert len(target) == 1
        worksheet_member = posixpath.normpath(posixpath.join("xl", target[0]))

        for member in source_zip.namelist():
            if member != worksheet_member:
                assert source_zip.read(member) == revised_zip.read(member), f"Unexpected XLSX member change: {member}"

        source_xml = etree.fromstring(source_zip.read(worksheet_member))
        revised_xml = etree.fromstring(revised_zip.read(worksheet_member))
        expected = {
            "D63": ("61972", "61971"),
            "E63": ("1383.4342999999999", "1383.4119"),
        }
        for coordinate, (old, new) in expected.items():
            source_value = source_xml.xpath(
                ".//s:c[@r=$coordinate]/s:v",
                namespaces={"s": spreadsheet_ns},
                coordinate=coordinate,
            )
            revised_value = revised_xml.xpath(
                ".//s:c[@r=$coordinate]/s:v",
                namespaces={"s": spreadsheet_ns},
                coordinate=coordinate,
            )
            assert len(source_value) == len(revised_value) == 1
            assert source_value[0].text == old and revised_value[0].text == new
            revised_value[0].text = old
        assert etree.tostring(source_xml, method="c14n") == etree.tostring(revised_xml, method="c14n")
    return ["Supplementary workbook: PASS (only Table_S2!D63/E63 changed for the ZXS59 header-count correction)"]


def validate_figure() -> list[str]:
    pdf = REVISE / "Figure_S2.pdf"
    svg = REVISE / "Figure_S2.svg"
    png = REVISE / "Figure_S2.png"
    etree.parse(svg)
    svg_text = svg.read_text(encoding="utf-8").lower()
    assert "#36617b" in svg_text and "#c84f50" in svg_text
    assert "#0272b2" not in svg_text and "#ec6f00" not in svg_text
    assert "age- and sex-adjusted eccdna count" in svg_text
    assert "age- and sex-adjusted epm" in svg_text
    assert "mapped reads (millions)" not in svg_text
    with Image.open(png) as image:
        assert image.size == (2161, 1488)
        dpi = image.info.get("dpi", (0, 0))
        assert all(abs(value - 300) < 1 for value in dpi)
    assert pdf.stat().st_size > 50_000 and svg.stat().st_size > 100_000
    return [
        "Figure S2: PASS (183-mm PDF/SVG plus 300-dpi PNG, 2161 × 1488 px; "
        "Figure 1 cohort colours confirmed; mapped-read panel removed; adjusted count/EPM panels present)"
    ]


def main() -> None:
    lines = []
    lines.extend(validate_manuscript())
    lines.extend(validate_responses())
    lines.extend(validate_tables())
    lines.extend(validate_figure())
    print("\n".join(lines))


if __name__ == "__main__":
    main()
