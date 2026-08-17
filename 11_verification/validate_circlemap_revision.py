#!/usr/bin/env python3
"""Validate Circle-Map robustness manuscript/revision outputs."""

from __future__ import annotations

import re
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_table_tools import cell_text, find_row, load_document as load_table_doc  # noqa: E402
from docx_tracked_edit import accepted_text, load_document  # noqa: E402


REVISE = Path(__file__).resolve().parent.parent
MAIN = REVISE / "Maintext_Covid-eccDNA_revised_tracked_v4.docx"
MAIN_BACKUP = REVISE / "Maintext_Covid-eccDNA_revised_tracked_v4.docx.pre_circlemap_bak"
WORKBOOK = REVISE / "Supplementary_Tables_revised.xlsx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def qn(tag: str) -> str:
    return f"{{{W}}}{tag}"


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def docx_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        assert_true(archive.testzip() is None, f"{path.name}: zip CRC failure")
        return archive.read("word/document.xml").decode("utf-8")


def check_docx_structure(path: Path) -> None:
    xml = docx_xml(path)
    root = ET.fromstring(xml)
    parents = {child: parent for parent in root.iter() for child in parent}

    def ancestors(node):
        current = parents.get(node)
        while current is not None:
            yield current
            current = parents.get(current)

    for node in root.iter(qn("delText")):
        assert_true(any(a.tag == qn("del") for a in ancestors(node)), f"{path.name}: w:delText outside w:del")
    for node in root.iter(qn("t")):
        assert_true(not any(a.tag == qn("del") for a in ancestors(node)), f"{path.name}: w:t inside w:del")
    assert_true("ns0:" not in xml and not re.search(r"\bns\d+:", xml), f"{path.name}: generated namespace prefix remains")


def count_endnote(path: Path) -> tuple[int, int]:
    xml = docx_xml(path)
    return xml.count("ADDIN EN.CITE"), xml.count("ADDIN EN.REFLIST")


def check_main_text() -> None:
    check_docx_structure(MAIN)
    assert_true(count_endnote(MAIN) == count_endnote(MAIN_BACKUP), "EndNote field counts changed")
    text = accepted_text(load_document(MAIN))
    required = [
        "All seven call sets supported higher normalized eccDNA abundance in COVID-19",
        "Median COVID-19/HC EPM ratios ranged from 3.02",
        "Circle_finder commit 3eb333db2ea6277dde36cbf640be9afeb710c717",
        "CTNNA2 as the highest-confidence computational target",
        "Supplementary Figure 6. Robustness of the COVID-19 versus HC eccDNA-burden conclusion",
        "Supplementary Figure 7. Robustness of downstream eccGene, chromatin and validation-target analyses",
        "Supplementary Table 12. Circle-Map call-set sensitivity analysis",
        "Supplementary Table 16. Validation-target support and strict breakpoint artifact-mask audit",
    ]
    for phrase in required:
        assert_true(phrase in text, f"main text missing phrase: {phrase}")
    stale = [
        "Use of a single caller may also propagate caller-specific errors",
        "Consensus calling with independent algorithms is identified as a priority for future validation",
    ]
    for phrase in stale:
        assert_true(phrase not in text, f"main text still contains stale phrase: {phrase}")


def check_responses() -> None:
    response_checks = [
        (
            REVISE / "JARE-D-26-04883_Reviewer_1_response.docx",
            "Potential methodological bias introduced by Circle-seq",
            "median COVID-19/HC EPM ratios ranged from 3.02 to 3.86",
        ),
        (
            REVISE / "JARE-D-26-04883_Reviewer_1_response.docx",
            "Lack of independent validation of disease-associated eccDNAs",
            "only CTNNA2 passed the strict breakpoint artifact mask",
        ),
        (
            REVISE / "JARE-D-26-04883_Reviewer_2_response.docx",
            "Lack of experimental validation and functional assessment",
            "computational robustness does not replace outward PCR/Sanger",
        ),
        (
            REVISE / "JARE-D-26-04883_Reviewer_3_response.docx",
            "Robustness of eccDNA identification and quantification",
            "we have now performed the requested caller/call-set robustness analysis",
        ),
        (
            REVISE / "JARE-D-26-04883_Reviewer_3_response.docx",
            "Definition and quantification of eccGene abundance",
            "Spearman ρ = 0.956, 0.956 and 0.955",
        ),
    ]
    for path, needle, expected in response_checks:
        check_docx_structure(path)
        tree = load_table_doc(path)
        _, _, _, cells = find_row(tree, needle, column=1)
        combined = cell_text(cells[2]) + "\n" + cell_text(cells[3])
        assert_true(expected in combined, f"{path.name}: missing response text for {needle}")


def check_workbook() -> None:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    expected = {
        "Table_S12": (600, 17),
        "Table_S13": (400, 15),
        "Table_S14": (2800, 11),
        "Table_S15": (300, 16),
        "Table_S16": (4500, 18),
    }
    for sheet, (min_rows, min_cols) in expected.items():
        assert_true(sheet in wb.sheetnames, f"workbook missing {sheet}")
        ws = wb[sheet]
        assert_true(ws.max_row >= min_rows, f"{sheet}: too few rows ({ws.max_row})")
        assert_true(ws.max_column >= min_cols, f"{sheet}: too few columns ({ws.max_column})")


def check_figures() -> None:
    for stem in ["Figure_S6", "Figure_S7"]:
        for suffix in ["pdf", "svg", "png", "caption.md", "palette.tsv", "style_QA.md"]:
            path = REVISE / f"{stem}.{suffix}" if suffix in {"pdf", "svg", "png"} else REVISE / f"{stem}_{suffix}"
            assert_true(path.exists(), f"missing {path.name}")
        subprocess.run(
            [
                sys.executable,
                str(REVISE / "scripts" / "check_figure_nature_spec.py"),
                str(REVISE / f"{stem}.pdf"),
                "--expect-width-mm",
                "183",
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
        )


def main() -> int:
    check_figures()
    check_workbook()
    check_main_text()
    check_responses()
    print("PASS: Circle-Map robustness revision outputs are internally consistent")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
