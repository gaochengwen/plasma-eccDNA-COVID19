#!/usr/bin/env python3
"""Build the age-matched pair assignments reported in Supplementary Table S19.

Matching rule (Reviewer 3, major comment 6)
-------------------------------------------
Each COVID-19 patient is matched 1:1 and without replacement to a healthy
control whose age differs by at most one year (caliper = 1 year). Patients with
no eligible control are left unmatched; no control is used twice. Among all
assignments satisfying the caliper the reported one is fixed by applying, in
order:

  1. the largest number of matched pairs,
  2. the smallest total absolute age difference over those pairs,
  3. the largest number of sex-concordant pairs,
  4. the lexicographically smallest list of (patient, control) identifiers.

Criteria 1-3 are optimised exactly with a linear-sum assignment; criterion 4 is
applied after enumerating every assignment that attains the same (1, 2, 3)
optimum, so the pairing is reproducible from the inputs alone. The enumeration
also yields the spread of the paired statistics across the tied optima, which is
what licenses the Methods statement about their invariance.

Inputs are the locked per-participant tables only: Supplementary Table S1 (age,
sex) and Supplementary Table S2a (detected eccDNA count, EPM). Nothing upstream
of those tables is recomputed here.
"""
from __future__ import annotations

import json
from itertools import combinations
from pathlib import Path

import numpy as np
import openpyxl
from scipy.optimize import linear_sum_assignment
from scipy.stats import wilcoxon

ROOT = Path("/home/gao/eccDNA")
XLSX = ROOT / "Submit" / "Supplementary_Tables_revised.xlsx"
OUT = ROOT / "revise" / "source_data"
CALIPER = 1
SEX_PENALTY = 0.001   # < 1/39 so age difference strictly outranks sex concordance
BIG = 1.0e6
SENSITIVITY_CALIPERS = (0, 1, 2, 3, 5)


def load():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    meta = {
        r[0]: {"rca": r[1], "sex": r[2], "age": int(r[3]), "group": r[4]}
        for r in wb["Table_S1"].iter_rows(min_row=3, values_only=True) if r and r[0]
    }
    burden = {
        r[0]: {"group": r[1], "mapped": r[2], "count": r[3], "epm": r[4]}
        for r in wb["Table_S2"].iter_rows(min_row=4, values_only=True)  # S2a block only
        if r and r[0] and r[1] in ("COVID-19", "HC") and isinstance(r[3], (int, float))
    }
    assert len(meta) == 78 and len(burden) == 78, (len(meta), len(burden))
    return meta, burden


def optimum(cov, hc, meta, caliper):
    """(n_pairs, total_age_diff, n_sex_concordant) of the optimal assignment."""
    cost = np.full((len(cov), len(hc)), BIG)
    for i, a in enumerate(cov):
        for j, b in enumerate(hc):
            d = abs(meta[a]["age"] - meta[b]["age"])
            if d <= caliper:
                cost[i, j] = d + (0.0 if meta[a]["sex"] == meta[b]["sex"] else SEX_PENALTY)
    ri, ci = linear_sum_assignment(cost)
    pairs = [(cov[i], hc[j]) for i, j in zip(ri, ci) if cost[i, j] < BIG / 2]
    tot = sum(abs(meta[a]["age"] - meta[b]["age"]) for a, b in pairs)
    sex = sum(1 for a, b in pairs if meta[a]["sex"] == meta[b]["sex"])
    return len(pairs), tot, sex


def _best(cov_sub, hc_sub, meta, caliper):
    """Optimal (n_pairs, total_age_diff, n_sex_concordant) on a sub-problem."""
    if not cov_sub or not hc_sub:
        return 0, 0, 0
    return optimum(cov_sub, hc_sub, meta, caliper)


def canonical_assignment(cov, hc, meta, caliper, target):
    """Lexicographically smallest assignment attaining `target`.

    Patients are visited in ascending identifier order and each is given the
    smallest-identifier control that still allows the optimum to be reached,
    which is checked by re-solving the assignment on what remains. Leaving a
    patient unmatched is only accepted when no control can be. This returns the
    same pairing as exhaustive enumeration but in polynomial time.
    """
    k_t, tot_t, sex_t = target
    order = sorted(cov)
    fixed: list[tuple[str, str]] = []
    used: set[str] = set()

    def reach(after_i, extra):
        k = len(fixed) + len(extra)
        tot = sum(abs(meta[a]["age"] - meta[b]["age"]) for a, b in fixed + extra)
        sex = sum(1 for a, b in fixed + extra if meta[a]["sex"] == meta[b]["sex"])
        rest_cov = order[after_i:]
        rest_hc = [b for b in hc if b not in used and all(b != y for _, y in extra)]
        dk, dt, ds = _best(rest_cov, rest_hc, meta, caliper)
        return (k + dk, tot + dt, sex + ds) == target

    for i, a in enumerate(order):
        cands = sorted(b for b in hc if b not in used
                       and abs(meta[a]["age"] - meta[b]["age"]) <= caliper)
        for b in cands:
            if reach(i + 1, [(a, b)]):
                fixed.append((a, b)); used.add(b)
                break
        else:
            if not reach(i + 1, []):
                raise RuntimeError(f"no optimal completion leaving {a} unmatched")
    assert len(fixed) == k_t
    return sorted(fixed)


def enumerate_optima(cov, hc, meta, caliper, target, cap=500000):
    """Every assignment attaining `target` = (n_pairs, total_diff, n_sex_conc)."""
    k_star, tot_star, sex_star = target
    cand = {a: [b for b in hc if abs(meta[a]["age"] - meta[b]["age"]) <= caliper] for a in cov}
    matchable = [a for a in cov if cand[a]]
    sols, truncated = [], [False]

    for chosen in combinations(matchable, k_star):
        rest = list(chosen)

        def rec(i, used, acc, tot, sex):
            if truncated[0]:
                return
            if tot > tot_star or (k_star - i) + sex < sex_star:
                return
            if i == len(rest):
                if tot == tot_star and sex == sex_star:
                    if len(sols) >= cap:
                        truncated[0] = True
                        return
                    sols.append(list(acc))
                return
            a = rest[i]
            for b in cand[a]:
                if b in used:
                    continue
                d = abs(meta[a]["age"] - meta[b]["age"])
                s = 1 if meta[a]["sex"] == meta[b]["sex"] else 0
                acc.append((a, b)); used.add(b)
                rec(i + 1, used, acc, tot + d, sex + s)
                used.discard(b); acc.pop()

        rec(0, set(), [], 0, 0)
    return sols, truncated[0]


def paired_stats(pairs, burden):
    c = np.array([[burden[a]["count"], burden[b]["count"]] for a, b in pairs], float)
    e = np.array([[burden[a]["epm"], burden[b]["epm"]] for a, b in pairs], float)
    return {
        "p_count": float(wilcoxon(c[:, 0] - c[:, 1], alternative="two-sided", mode="exact").pvalue),
        "p_epm": float(wilcoxon(e[:, 0] - e[:, 1], alternative="two-sided", mode="exact").pvalue),
        "fc_count": float(np.median(c[:, 0] / c[:, 1])),
        "fc_epm": float(np.median(e[:, 0] / e[:, 1])),
        "n_count_up": int((c[:, 0] > c[:, 1]).sum()),
        "n_epm_up": int((e[:, 0] > e[:, 1]).sum()),
    }


def main():
    meta, burden = load()
    cov = sorted([k for k, v in meta.items() if v["group"] == "COVID-19"],
                 key=lambda k: (meta[k]["age"], k))
    hc = sorted([k for k, v in meta.items() if v["group"] == "HC"],
                key=lambda k: (meta[k]["age"], k))

    target = optimum(cov, hc, meta, CALIPER)
    pairs = canonical_assignment(cov, hc, meta, CALIPER, target)
    # exhaustive check at the primary caliper: confirms the greedy selection and
    # measures how far the paired statistics move across the tied optima
    sols, truncated = enumerate_optima(cov, hc, meta, CALIPER, target)
    assert sols and not truncated, (len(sols), truncated)
    sols = [sorted(s) for s in sols]
    assert pairs == min(sols), "canonical_assignment disagrees with enumeration"
    pairs.sort(key=lambda p: (meta[p[0]]["age"], p[0]))

    st = paired_stats(pairs, burden)
    across = [paired_stats(s, burden) for s in sols]
    prov = {
        "caliper_years": CALIPER,
        "rule": ("1:1 nearest-neighbour matching without replacement within a 1-year age caliper; "
                 "maximise pairs, then minimise total absolute age difference, then maximise sex "
                 "concordance, then take the lexicographically smallest identifier list"),
        "n_pairs": len(pairs),
        "n_tied_optimal_assignments": len(sols),
        "total_age_difference_years": target[1],
        "max_pair_age_difference_years": max(abs(meta[a]["age"] - meta[b]["age"]) for a, b in pairs),
        "n_sex_concordant_pairs": target[2],
        "mean_age_covid": float(np.mean([meta[a]["age"] for a, _ in pairs])),
        "mean_age_hc": float(np.mean([meta[b]["age"] for _, b in pairs])),
        "n_covid_unmatched": len(cov) - len(pairs),
        "female_patients_matched": sum(1 for a, _ in pairs if meta[a]["sex"] == "Female"),
        "female_controls_matched": sum(1 for _, b in pairs if meta[b]["sex"] == "Female"),
        **{f"{k}": v for k, v in st.items()},
        "p_count_across_optima": sorted({round(x["p_count"], 8) for x in across}),
        "p_epm_across_optima": sorted({round(x["p_epm"], 8) for x in across}),
        "fc_count_range": [min(x["fc_count"] for x in across), max(x["fc_count"] for x in across)],
        "fc_epm_range": [min(x["fc_epm"] for x in across), max(x["fc_epm"] for x in across)],
        "inputs": {"ages_and_sex": "Supplementary Table S1", "count_and_epm": "Supplementary Table S2a"},
    }

    sens = []
    for cal in SENSITIVITY_CALIPERS:
        t = optimum(cov, hc, meta, cal)
        pp = canonical_assignment(cov, hc, meta, cal, t)
        s = paired_stats(pp, burden)
        sens.append({"caliper_years": cal, "n_pairs": t[0],
                     "total_age_difference_years": t[1],
                     "max_pair_age_difference_years": max(abs(meta[a]["age"] - meta[b]["age"]) for a, b in pp),
                     "n_sex_concordant_pairs": t[2], **s})

    OUT.mkdir(parents=True, exist_ok=True)
    rows = [[a, meta[a]["age"], meta[a]["sex"], burden[a]["count"], round(burden[a]["epm"], 4),
             b, meta[b]["age"], meta[b]["sex"], burden[b]["count"], round(burden[b]["epm"], 4),
             abs(meta[a]["age"] - meta[b]["age"]),
             round(burden[a]["count"] / burden[b]["count"], 3),
             round(burden[a]["epm"] / burden[b]["epm"], 3)] for a, b in pairs]
    (OUT / "age_matched_pairs.json").write_text(json.dumps(rows, indent=1))
    (OUT / "age_matched_pairs_provenance.json").write_text(json.dumps(prov, indent=1))
    (OUT / "age_matched_caliper_sensitivity.json").write_text(json.dumps(sens, indent=1))

    print(json.dumps(prov, indent=1))
    print("\ncaliper sensitivity:")
    for s in sens:
        print("  {caliper_years}y  n={n_pairs:2d} totalD={total_age_difference_years:2d} "
              "count P={p_count:.5f} ({n_count_up}/{n_pairs}) FC={fc_count:.2f}  "
              "EPM P={p_epm:.5f} ({n_epm_up}/{n_pairs}) FC={fc_epm:.2f}".format(**s))


if __name__ == "__main__":
    main()
