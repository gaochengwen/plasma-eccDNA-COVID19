#!/usr/bin/env python3
"""Re-derive the age-matched pairing and Table S4 on the v2 call set.

`revise/scripts/build_age_matched_pairs.py` reads the per-participant burden
from Table S2a of the v1 submission workbook. Table S2a was never migrated, so
its whole output chain -- the canonical pair file, the caliper sensitivity
table, Table S4 and the Methods sentence about the tied optima -- stayed on the
v1 call set even though the Results text was updated to v2.

This runs the original matching implementation unchanged, against the repaired
v2 workbook, so the pairing rule and the tie enumeration are byte-for-byte the
same logic that produced the submitted numbers. Only the burden inputs differ.

The matching itself is on age and sex alone, so the 12 pairs are expected to be
identical; what changes is every statistic computed from them.

Must run after `repair_v1_supplementary_tables_20260813.py`.
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import numpy as np
import openpyxl

RUN = Path(__file__).resolve().parents[2]
PROJECT = RUN.parent
WB = RUN / "tables" / "supplementary" / "Supplementary_Tables_v2.xlsx"
OUT = RUN / "tables" / "age_matching"
UPSTREAM = PROJECT / "revise" / "scripts" / "build_age_matched_pairs.py"


def load_upstream():
    spec = importlib.util.spec_from_file_location("build_age_matched_pairs", UPSTREAM)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def load_v2(m):
    """m.load(), but against the repaired v2 workbook."""
    wb = openpyxl.load_workbook(WB, read_only=True, data_only=True)
    meta = {
        r[0]: {"rca": r[1], "sex": r[2], "age": int(r[3]), "group": r[4]}
        for r in wb["Table_S1"].iter_rows(min_row=3, values_only=True) if r and r[0]
    }
    burden = {
        r[0]: {"group": r[1], "mapped": r[2], "count": r[3], "epm": r[4]}
        for r in wb["Table_S2"].iter_rows(min_row=4, values_only=True)
        if r and r[0] and r[1] in ("COVID-19", "HC") and isinstance(r[3], (int, float))
    }
    assert len(meta) == 78 and len(burden) == 78, (len(meta), len(burden))
    total = sum(v["count"] for v in burden.values())
    if total != 10_099_219:
        raise SystemExit(f"Table S2a still not on v2: burden sums to {total:,}")
    return meta, burden


def main() -> int:
    m = load_upstream()
    meta, burden = load_v2(m)
    cov = sorted([k for k, v in meta.items() if v["group"] == "COVID-19"],
                 key=lambda k: (meta[k]["age"], k))
    hc = sorted([k for k, v in meta.items() if v["group"] == "HC"],
                key=lambda k: (meta[k]["age"], k))

    target = m.optimum(cov, hc, meta, m.CALIPER)
    pairs = m.canonical_assignment(cov, hc, meta, m.CALIPER, target)
    sols, truncated = m.enumerate_optima(cov, hc, meta, m.CALIPER, target)
    assert sols and not truncated, (len(sols), truncated)
    sols = [sorted(s) for s in sols]
    assert pairs == min(sols), "canonical_assignment disagrees with enumeration"
    pairs.sort(key=lambda p: (meta[p[0]]["age"], p[0]))

    st = m.paired_stats(pairs, burden)
    across = [m.paired_stats(s, burden) for s in sols]
    prov = {
        "call_set": "circlemap_methods (v2)",
        "caliper_years": m.CALIPER,
        "n_pairs": len(pairs),
        "n_tied_optimal_assignments": len(sols),
        "total_age_difference_years": target[1],
        "max_pair_age_difference_years": max(abs(meta[a]["age"] - meta[b]["age"])
                                             for a, b in pairs),
        "n_sex_concordant_pairs": target[2],
        "mean_age_covid": float(np.mean([meta[a]["age"] for a, _ in pairs])),
        "mean_age_hc": float(np.mean([meta[b]["age"] for _, b in pairs])),
        **st,
        "p_count_across_optima": sorted({round(x["p_count"], 10) for x in across}),
        "p_epm_across_optima": sorted({round(x["p_epm"], 10) for x in across}),
        "fc_count_range": [min(x["fc_count"] for x in across),
                           max(x["fc_count"] for x in across)],
        "fc_epm_range": [min(x["fc_epm"] for x in across),
                         max(x["fc_epm"] for x in across)],
        "inputs": {"ages_and_sex": "Supplementary Table S1",
                   "count_and_epm": "Supplementary Table S2a (v2)"},
    }

    sens = []
    for cal in m.SENSITIVITY_CALIPERS:
        t = m.optimum(cov, hc, meta, cal)
        pp = m.canonical_assignment(cov, hc, meta, cal, t)
        s = m.paired_stats(pp, burden)
        sens.append({"caliper_years": cal, "n_pairs": t[0],
                     "total_age_difference_years": t[1],
                     "max_pair_age_difference_years": max(
                         abs(meta[a]["age"] - meta[b]["age"]) for a, b in pp),
                     "n_sex_concordant_pairs": t[2], **s})

    rows = [[a, meta[a]["age"], meta[a]["sex"], burden[a]["count"],
             round(burden[a]["epm"], 4),
             b, meta[b]["age"], meta[b]["sex"], burden[b]["count"],
             round(burden[b]["epm"], 4),
             abs(meta[a]["age"] - meta[b]["age"]),
             round(burden[a]["count"] / burden[b]["count"], 3),
             round(burden[a]["epm"] / burden[b]["epm"], 3)] for a, b in pairs]

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "age_matched_pairs_v2.json").write_text(json.dumps(rows, indent=1))
    (OUT / "age_matched_pairs_provenance_v2.json").write_text(json.dumps(prov, indent=1))
    (OUT / "age_matched_caliper_sensitivity_v2.json").write_text(json.dumps(sens, indent=1))

    # ---------- write Table S4 -------------------------------------------
    wb = openpyxl.load_workbook(WB)
    ws = wb["Table_S4"]
    for i, row in enumerate(rows):                      # S4a, header row 3
        for c, v in enumerate(row, start=1):
            ws.cell(row=4 + i, column=c).value = v
    keys = ["caliper_years", "n_pairs", "total_age_difference_years",
            "max_pair_age_difference_years", "n_sex_concordant_pairs",
            "p_count", "n_count_up", "fc_count", "p_epm", "n_epm_up", "fc_epm"]
    for i, s in enumerate(sens):                        # S4b, header row 18
        for c, k in enumerate(keys, start=1):
            ws.cell(row=19 + i, column=c).value = s[k]
    wb.save(WB)

    print(json.dumps(prov, indent=1))
    print("\ncaliper sensitivity (v2):")
    for s in sens:
        print("  {caliper_years}y n={n_pairs:2d} totalD={total_age_difference_years:2d} "
              "count P={p_count:.6g} ({n_count_up}/{n_pairs}) FC={fc_count:.3f} | "
              "EPM P={p_epm:.6g} ({n_epm_up}/{n_pairs}) FC={fc_epm:.3f}".format(**s))
    print(f"\nsaved {OUT} and Table_S4 in {WB.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
