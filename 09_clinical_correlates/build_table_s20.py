#!/usr/bin/env python3
"""Build Supplementary Table S21: the non-COVID pneumonia comparator samples.

These nine participants were collected and clinically annotated under the same
approval as the sequenced cohort but were never taken into the laboratory
workflow: there is no post-RCA yield, no library, no reads and no eccDNA call
for any of them. They therefore contribute no data to any analysis in this
study, and the table exists to declare their existence rather than to support a
comparison.

Diagnoses are the sampling-time diagnoses supplied by the clinical team, not
the `main_diagnosis` field of the index visit -- for inpatients that field
carries the admitting department's primary diagnosis (haematology coded the
leukaemia, nephrology coded the renal disease) and understates the pneumonia.

Patient identifiers are deliberately limited to the internal sample ID; the
hospital admission numbers used to link the records are not published, matching
Supplementary Table S1.

Run from the repository root.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[3]
XLSX = ROOT / "Submit" / "Supplementary_Tables_revised.xlsx"
TRACE = ROOT / "临床信息" / "样本信息_溯源表.csv"

# sampling-time diagnosis supplied by the clinical team, keyed by sample ID
DIAGNOSIS = {
    "NCP7440": "LAE pneumonia",
    "NCP6694": "Aspiration pneumonia",
    "NCP16486": "LAE pneumonia",
    "NCP13762": "LAE pneumonia",
    "NCP12581": "LAE pneumonia",
    "NCPs862": "LAE pneumonia",
    "NCP322": "Interstitial pneumonia",
    "NCP307": "Interstitial pneumonia",
    "NCP147": "Interstitial pneumonia",
}
ORDER = ["NCP7440", "NCP6694", "NCP16486", "NCP13762", "NCP12581",
         "NCPs862", "NCP322", "NCP307", "NCP147"]

HEADERS = ["Sample ID", "Gender", "Age", "Diagnosis at sampling",
           "Days from diagnosis to sampling", "Treatment before sampling",
           "Severity at sampling", "ICU at sampling",
           "Respiratory support at sampling", "Comorbidities",
           "CRP (mg/L)", "D-dimer (ng/mL)",
           "Lymphocyte count (10^9/L)", "Clinical outcome"]

TITLE = ("Supplementary Table S21. Non-COVID pneumonia samples collected but not "
         "sequenced in this study (n = 9)")

NOTE = (
    "These nine participants were recruited and clinically annotated under the same "
    "ethics approval (QYFYEC2024-187) as the sequenced cohort, but none entered the "
    "Circle-seq workflow: no post-rolling-circle-amplification product, sequencing "
    "library, read or eccDNA call exists for any of them, and they contribute to no "
    "analysis reported in this study. The table is provided so that their existence "
    "is on the record rather than to support any comparison. Diagnoses are the "
    "sampling-time diagnoses recorded by the clinical team; note that for the "
    "inpatients the admitting department's primary diagnosis is not the pneumonia "
    "(for example the haematology admissions are coded to the underlying "
    "haematological disease). Clinical variables follow the same sampling-anchored "
    "definitions as Supplementary Table S1. IL-6 and ferritin were not measured in "
    "any of these participants, and no participant died in hospital. These samples "
    "were collected between February 2025 "
    "and March 2026, whereas the COVID-19 cohort was sampled in January and February "
    "2023, so sampling batch is completely confounded with group; any future "
    "comparison would have to be designed around that, for example by processing "
    "both groups in a common batch."
)


def main() -> None:
    trace = pd.read_csv(TRACE).set_index("Sample ID")
    filled = pd.read_excel(ROOT / "临床信息" / "样本信息_filled.xlsx", header=1)
    filled = filled.set_index("Sample ID")

    shutil.copy2(XLSX, XLSX.with_suffix(".xlsx.pre_s21"))
    wb = openpyxl.load_workbook(XLSX)
    if "Table_S21" in wb.sheetnames:
        del wb["Table_S21"]
    ws = wb.create_sheet("Table_S21")

    ws.cell(1, 1).value = TITLE
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(2, c)
        cell.value = h
        # marks the header row for the shared formatter's role classifier
        cell.font = openpyxl.styles.Font(bold=True)

    def clean(v, blank="Not documented"):
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return blank
        s = str(v).strip()
        return blank if s.lower() in ("nan", "") else s

    for i, sid in enumerate(ORDER):
        t = trace.loc[sid]
        f = filled.loc[sid]
        row = [
            sid,
            clean(f.get("Gender")),
            int(f.get("Age")),
            DIAGNOSIS[sid],
            clean(t.get("Days_Diagnosis_to_Sampling"), "Not applicable"),
            clean(t.get("Treatment_Before_Sampling")),
            clean(t.get("Severity")),
            clean(t.get("Ward_ICU")),
            clean(t.get("Respiratory_Support")),
            clean(t.get("Comorbidities")),
            clean(t.get("CRP"), "Not measured"),
            clean(t.get("D_dimer"), "Not measured"),
            clean(t.get("Lymphocyte_Count"), "Not measured"),
            clean(t.get("Outcome")),
        ]
        for c, v in enumerate(row, start=1):
            ws.cell(3 + i, c).value = v

    note_row = 3 + len(ORDER) + 1
    ws.cell(note_row, 1).value = NOTE

    last = openpyxl.utils.get_column_letter(len(HEADERS))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A{note_row}:{last}{note_row}")

    wb.save(XLSX)

    print(f"wrote Table_S21: {len(ORDER)} participants, {len(HEADERS)} columns")
    for sid in ORDER:
        print(f"  {sid:<10} {DIAGNOSIS[sid]:<24} "
              f"age {int(filled.loc[sid,'Age'])} {filled.loc[sid,'Gender']}")


if __name__ == "__main__":
    main()
