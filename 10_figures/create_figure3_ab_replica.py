#!/usr/bin/env python3
"""Recreate Figure 3A-B with revised eccGene data and the original layout."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.collections import LineCollection, PatchCollection
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import Rectangle


ROOT = Path(__file__).resolve().parents[2]
ANALYSIS = ROOT / "analyse" / "eccGene"
RESULTS = ANALYSIS / "results"
MATRICES = ANALYSIS / "matrices"
REVISE = ROOT / "revise"
SOURCE_DATA = REVISE / "source_data"

MM = 1 / 25.4
COVID_RED = "#CF3F42"
HC_TEAL = "#2F627A"
DETECTED_BLUE = "#3F73BF"
NOT_DETECTED_GREY = "#D2D2D2"
LOW_BLUE = "#3D70B3"
MID_NEUTRAL = "#F7F7F7"
HIGH_RED = "#E1251B"


plt.rcParams.update(
    {
        "font.family": "Arial",
        "font.size": 6,
        "axes.linewidth": 0.6,
        "xtick.major.width": 0.5,
        "ytick.major.width": 0.5,
        "xtick.major.size": 2,
        "ytick.major.size": 2,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "svg.fonttype": "none",
        "savefig.facecolor": "white",
    }
)


def read_tsv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, sep="\t")


def sample_order() -> tuple[list[str], list[str]]:
    metadata = read_tsv(ANALYSIS / "sample_metadata.tsv")
    group = metadata.set_index("sample_id")["group"].to_dict()

    # Preserve the within-group ordering used in the original Table S7/Figure 3.
    old_workbook = ROOT / "Maintext" / "Supplementary_Tables.xlsx"
    from openpyxl import load_workbook

    workbook = load_workbook(old_workbook, read_only=True, data_only=True)
    header = list(
        next(
            workbook["Table_S7"].iter_rows(
                min_row=2,
                max_row=2,
                values_only=True,
            )
        )
    )
    workbook.close()
    old_samples = [str(value) for value in header[1:-3] if value is not None]
    covid = [sample for sample in old_samples if group[sample] == "COVID-19"]
    hc = [sample for sample in old_samples if group[sample] == "HC"]
    if len(covid) != 39 or len(hc) != 39:
        raise ValueError(f"Unexpected group sizes: COVID={len(covid)}, HC={len(hc)}")
    return covid, hc


def prepare_panel_a(
    ordered_samples: list[str],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    statistics = read_tsv(RESULTS / "junction_abundance_wilcoxon.tsv")
    selected = statistics.loc[
        (statistics["wilcoxon_q_BH"] < 0.05)
        & (statistics["log2FC_mean_EA_COVID_vs_HC"].abs() >= 1)
    ].copy()
    selected = selected.sort_values(
        [
            "direction",
            "wilcoxon_q_BH",
            "log2FC_mean_EA_COVID_vs_HC",
            "Gene",
        ],
        ascending=[False, True, False, True],
        kind="stable",
    )
    if len(selected) != 3402:
        raise ValueError(f"Expected 3,402 abundance genes, observed {len(selected)}")

    ea = read_tsv(MATRICES / "junction_EA.tsv").set_index("Gene")
    values = ea.loc[selected["Gene"], ordered_samples].astype(float)
    transformed = np.log1p(values.to_numpy())
    means = transformed.mean(axis=1, keepdims=True)
    standard_deviations = transformed.std(axis=1, ddof=0, keepdims=True)
    standard_deviations[standard_deviations == 0] = 1
    z_scores = np.clip(
        (transformed - means) / standard_deviations,
        -2.5,
        2.5,
    )
    z_frame = pd.DataFrame(
        z_scores,
        index=selected["Gene"],
        columns=ordered_samples,
    )
    export_columns = [
        "Gene",
        "mean_EA_COVID",
        "mean_EA_HC",
        "log2FC_mean_EA_COVID_vs_HC",
        "cliffs_delta_COVID_vs_HC",
        "wilcoxon_p",
        "wilcoxon_q_BH",
        "direction",
    ]
    selected[export_columns].to_csv(
        SOURCE_DATA / "Figure_3A_replica_abundance_genes.tsv",
        sep="\t",
        index=False,
    )
    z_frame.reset_index().to_csv(
        SOURCE_DATA / "Figure_3A_replica_log1p_EA_row_zscore.tsv",
        sep="\t",
        index=False,
    )
    return z_frame, selected


def prepare_panel_b(
    abundance_genes: set[str],
    ordered_samples: list[str],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    detection = read_tsv(RESULTS / "junction_detection_fisher.tsv")
    selected = detection.loc[
        detection["Gene"].isin(abundance_genes)
        & (detection["fisher_q_BH"] < 0.05)
        & (detection["hc_detected"] == 0)
        & (detection["frequency_difference_COVID_minus_HC"] > 0)
    ].copy()
    # Preserve the documented top-30 selection, then order the displayed genes
    # by recurrence (COVID-19 samples detected).  Fisher q and gene symbol are
    # deterministic tie-breakers.
    selected = selected.sort_values(
        ["fisher_q_BH", "covid_detected", "Gene"],
        ascending=[True, False, True],
        kind="stable",
    ).head(30)
    selected = selected.sort_values(
        ["covid_detected", "fisher_q_BH", "Gene"],
        ascending=[False, True, True],
        kind="stable",
    )
    if len(selected) != 30:
        raise ValueError(f"Expected 30 detection genes, observed {len(selected)}")

    detected = read_tsv(MATRICES / "junction_detected.tsv").set_index("Gene")
    matrix = detected.loc[selected["Gene"], ordered_samples].astype(int)
    covid_samples = ordered_samples[:39]
    hc_samples = ordered_samples[39:]
    covid_recurrence = matrix[covid_samples].sum(axis=0)
    covid_samples = sorted(
        covid_samples,
        key=lambda sample: (-int(covid_recurrence[sample]), sample),
    )
    matrix = matrix.loc[:, covid_samples + hc_samples]
    selected.to_csv(
        SOURCE_DATA / "Figure_3B_replica_detection_genes.tsv",
        sep="\t",
        index=False,
    )
    matrix.reset_index().to_csv(
        SOURCE_DATA / "Figure_3B_replica_detection_matrix.tsv",
        sep="\t",
        index=False,
    )
    return matrix, selected


def add_group_strip(
    ax: plt.Axes,
    covid_fraction: float,
    y: float,
    height: float,
) -> None:
    ax.add_patch(
        Rectangle(
            (0, y),
            covid_fraction,
            height,
            transform=ax.transAxes,
            facecolor=COVID_RED,
            edgecolor="none",
            clip_on=False,
        )
    )
    ax.add_patch(
        Rectangle(
            (covid_fraction, y),
            1 - covid_fraction,
            height,
            transform=ax.transAxes,
            facecolor=HC_TEAL,
            edgecolor="none",
            clip_on=False,
        )
    )


def plot_replica(
    z_frame: pd.DataFrame,
    abundance: pd.DataFrame,
    detection_matrix: pd.DataFrame,
    detection: pd.DataFrame,
    covid_samples: list[str],
    hc_samples: list[str],
) -> None:
    figure = plt.figure(figsize=(183 * MM, 78 * MM), facecolor="white")

    # Panel A geometry mirrors the narrow, high heatmap in the original figure.
    ax_a = figure.add_axes([0.045, 0.19, 0.258, 0.675])
    abundance_cmap = LinearSegmentedColormap.from_list(
        "original_abundance",
        [
            (0.00, LOW_BLUE),
            (0.35, "#9EBBD8"),
            (0.50, MID_NEUTRAL),
            (0.65, "#F2B1A8"),
            (1.00, HIGH_RED),
        ],
    )
    image_a = ax_a.imshow(
        z_frame.to_numpy(),
        aspect="auto",
        interpolation="nearest",
        cmap=abundance_cmap,
        vmin=-2.5,
        vmax=2.5,
        rasterized=True,
    )
    ax_a.set_xticks([])
    ax_a.set_yticks([])
    for spine in ax_a.spines.values():
        spine.set_visible(False)
    add_group_strip(ax_a, 0.5, y=-0.037, height=0.016)
    ax_a.text(
        0.25,
        -0.068,
        "COVID-19",
        transform=ax_a.transAxes,
        ha="center",
        va="top",
        fontsize=8,
    )
    ax_a.text(
        0.75,
        -0.068,
        "HC",
        transform=ax_a.transAxes,
        ha="center",
        va="top",
        fontsize=8,
    )

    color_ax = figure.add_axes([0.149, 0.902, 0.118, 0.035])
    colorbar = figure.colorbar(image_a, cax=color_ax, orientation="horizontal")
    colorbar.set_ticks([])
    colorbar.outline.set_visible(False)
    figure.text(0.073, 0.918, "eccDNA\nabundance", ha="left", va="center", fontsize=6)
    figure.text(0.149, 0.890, "Low", ha="center", va="top", fontsize=6)
    figure.text(0.267, 0.890, "High", ha="center", va="top", fontsize=6)
    figure.text(0.014, 0.948, "A", fontsize=13, fontweight="bold", va="top")

    # Panel B retains the original binary matrix, top sample totals, and right labels.
    ax_b = figure.add_axes([0.365, 0.19, 0.500, 0.615])
    rows, columns = detection_matrix.shape
    ax_b.add_patch(
        Rectangle(
            (0, 0),
            columns,
            rows,
            facecolor=NOT_DETECTED_GREY,
            edgecolor="none",
        )
    )
    detected_cells = [
        Rectangle((column, row), 1, 1)
        for row, column in np.argwhere(detection_matrix.to_numpy() == 1)
    ]
    ax_b.add_collection(
        PatchCollection(
            detected_cells,
            facecolor=DETECTED_BLUE,
            edgecolor="none",
            antialiased=False,
        )
    )
    vertical_lines = [
        [(column, 0), (column, rows)] for column in range(columns + 1)
    ]
    horizontal_lines = [
        [(0, row), (columns, row)] for row in range(rows + 1)
    ]
    ax_b.add_collection(
        LineCollection(
            vertical_lines + horizontal_lines,
            colors="white",
            linewidths=0.55,
            antialiaseds=False,
        )
    )
    ax_b.set_xlim(0, columns)
    ax_b.set_ylim(rows, 0)
    ax_b.set_xticks([])
    ax_b.set_yticks([])
    for spine in ax_b.spines.values():
        spine.set_visible(False)
    add_group_strip(ax_b, 0.5, y=-0.037, height=0.016)
    ax_b.text(
        0.25,
        -0.068,
        "COVID-19",
        transform=ax_b.transAxes,
        ha="center",
        va="top",
        fontsize=8,
    )
    ax_b.text(
        0.75,
        -0.068,
        "HC",
        transform=ax_b.transAxes,
        ha="center",
        va="top",
        fontsize=8,
    )
    for index, gene in enumerate(detection_matrix.index):
        ax_b.text(
            columns + 1.15,
            index + 0.5,
            gene,
            ha="left",
            va="center",
            fontsize=5.0,
            fontstyle="italic",
            clip_on=False,
        )
    ax_b.text(
        columns + 1.15,
        -1.05,
        "Genes",
        ha="left",
        va="bottom",
        fontsize=6,
        clip_on=False,
    )

    bar_ax = figure.add_axes([0.365, 0.825, 0.500, 0.115], sharex=ax_b)
    panel_b_covid_samples = detection_matrix.columns[: len(covid_samples)]
    covid_totals = (
        detection_matrix.loc[:, panel_b_covid_samples].sum(axis=0).to_numpy()
    )
    bar_ax.bar(
        np.arange(len(covid_samples)) + 0.5,
        covid_totals,
        width=0.78,
        color=DETECTED_BLUE,
        edgecolor="none",
    )
    bar_ax.set_xlim(0, columns)
    bar_ax.set_ylim(0, 30)
    bar_ax.set_xticks([])
    bar_ax.set_yticks([0, 10, 20, 30])
    bar_ax.tick_params(axis="y", labelsize=5, pad=1)
    bar_ax.spines["top"].set_visible(False)
    bar_ax.spines["right"].set_visible(False)
    bar_ax.spines["bottom"].set_visible(False)
    bar_ax.spines["left"].set_linewidth(0.5)
    figure.text(0.326, 0.948, "B", fontsize=13, fontweight="bold", va="top")

    for stem, kwargs in (
        ("Figure_3_AB", {"format": "pdf", "dpi": 600}),
        ("Figure_3_AB", {"format": "svg", "dpi": 600}),
        ("Figure_3_AB", {"format": "png", "dpi": 300}),
    ):
        figure.savefig(
            REVISE / f"{stem}.{kwargs['format']}",
            **kwargs,
            bbox_inches=None,
            facecolor="white",
        )
    plt.close(figure)

    metadata = {
        "reference_figure": str(ROOT / "Maintext" / "Figure 3.pdf"),
        "main_definition": "junction",
        "panel_a_filter": "BH-FDR < 0.05 and absolute mean-EA log2FC >= 1",
        "panel_a_gene_count": int(len(abundance)),
        "panel_a_transform": "row z-score of log1p(EA), clipped to [-2.5, 2.5]",
        "panel_b_filter": (
            "Panel A genes with detection BH-FDR < 0.05, HC detected count = 0, "
            "positive COVID-HC frequency difference; top 30 selected by q, COVID count, "
            "and gene, then displayed by decreasing COVID recurrence, q, and gene"
        ),
        "panel_b_sample_order": (
            "COVID samples by decreasing number of the 30 displayed genes detected, "
            "then sample ID; HC samples retain the reference order"
        ),
        "panel_b_gene_count": int(len(detection)),
        "sample_count": int(len(covid_samples) + len(hc_samples)),
        "covid_samples": int(len(covid_samples)),
        "hc_samples": int(len(hc_samples)),
        "output_size_mm": [183, 78],
    }
    (SOURCE_DATA / "Figure_3_AB_replica_metadata.json").write_text(
        json.dumps(metadata, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    SOURCE_DATA.mkdir(parents=True, exist_ok=True)
    covid_samples, hc_samples = sample_order()
    ordered_samples = covid_samples + hc_samples
    z_frame, abundance = prepare_panel_a(ordered_samples)
    detection_matrix, detection = prepare_panel_b(
        set(abundance["Gene"]),
        ordered_samples,
    )
    plot_replica(
        z_frame,
        abundance,
        detection_matrix,
        detection,
        covid_samples,
        hc_samples,
    )
    print("Created revised Figure 3A-B replica.")


if __name__ == "__main__":
    main()
